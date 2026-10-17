"""信息图式科研页（info 风格，内容页默认）。

视觉规范取自用户给的《蓝白科研答辩模板｜12 页》（Final_editable.pptx 实测）与“像素答辩”参考页：
白底、深蓝 #073DA5 标题与标题带、细蓝灰边框 #9CB3D3（0.6pt）、浅蓝 #EAF0FC 小标题条与说明块、
黑色 #111722 正文、少量红色 #DE1515 强调、底部 ■ 结论条。字号用模板实测值（标题 32、引导句 16.8、
标题带 20.4、小标题 15.6、正文 13.2–15），不靠放大字号填空白，空白用更多具体内容填。

页面写法：
    {"style": "info", "title": "2.1 两种方案对比",
     "lead": {"text": "…", "highlights": ["…"], "tag": "核心目标"},   # ✓ 引导句（tag 换成深蓝标签）
     "sublead": {"text": "…"},                                         # 无框居中深蓝副引导句（可省）
     "band": "方法对比：从原理、性能到应用选择",                         # 全宽标题带；{"text", "images"} 为图片衬底
     "process": ["数据标准化", "算法自主化", …],                         # 全宽流程条（可省）
     "layout": "info_split", "regions": {"left": {...}, "right": {...}},
     "summary": {"text": "结论与启示：…", "highlights": ["…"]}}
region 也可以不用预设槽位，直接写 "grid": [起始列, 列数, 起始行, 行数] 自由拼版。
"""
import math
import re
from pathlib import Path
from PIL import Image
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from styled_deck import slide_number, zh_run
from textfit import measure, text_width
import layouts

ASSETS = Path(__file__).resolve().parents[1] / 'assets'

# ---------------------------------------------------------------- 颜色（模板实测）
NAVY = '073DA5'       # 标题、标题带、标签、编号、方块符号
INK = '111722'        # 正文
RED = 'DE1515'        # 强调
LINE = '9CB3D3'       # 分区细边框
LIGHT = 'EAF0FC'      # 小标题条、说明块、单元格
PALE = 'F8FAFE'       # 图位底色
ARROW = 'BCD3F6'      # 浅蓝大箭头
RULE = 'D5DEEB'       # 页眉下分隔线
GREY = '6F7D90'       # 时间、占位说明
DIVIDER = '073DA5'
LAYER = ['073DA5', '1F56B8', '3E73CB', '6A93DA', '9DB8E8']   # 由深到浅

# ---------------------------------------------------------------- 几何（英寸，16:9 = 13.333 × 7.5）
X0, W0 = 0.32, 12.7
LEAD_Y = 0.94
BAND_H = 0.49
SUMMARY_H = 0.53
BOTTOM = 7.5 - 0.2
GAP_Y = 0.1

# ---------------------------------------------------------------- 字号（pt，模板实测；放不下才往下收到 12）
SIZE = {'title': 32, 'num': 33, 'lead': 16.8, 'check': 20.4, 'band': 20.4, 'sublead': 15.6,
        'summary': 15.6, 'region_band': 18, 'head': 15.6, 'note': 15, 'markers': 14.4, 'chevrons': 14.4,
        'para': 15, 'hexes': 14.4, 'pill': 15.6, 'process': 15.6, 'side': 15.6, 'label': 15.6,
        'foot': 14.4, 'caption': 13.2, 'table': 13.2, 'cell': 15.6, 'small': 12}
AUTOFIT_MIN = 12

# ---------------------------------------------------------------- 版式预设（24 列 × 12 行网格）
INFO_LAYOUTS = {
    # 一个主体
    'info_full': {'main': [0, 24, 0, 12]},
    'info_focus': {'main': [0, 24, 0, 9], 'bottom': [0, 24, 9, 3]},
    'info_hero': {'main': [0, 14, 0, 12], 'side': [14, 10, 0, 12]},
    'info_hero_left': {'side': [0, 10, 0, 12], 'main': [10, 14, 0, 12]},
    'info_side3': {'main': [0, 14, 0, 12], 's1': [14, 10, 0, 4], 's2': [14, 10, 4, 4], 's3': [14, 10, 8, 4]},
    'info_big': {'main': [0, 17, 0, 12], 'side': [17, 7, 0, 12]},            # 大图页：图占 70%，不放标题带
    'info_big_left': {'side': [0, 7, 0, 12], 'main': [7, 17, 0, 12]},
    # 并列
    'info_split': {'left': [0, 12, 0, 12], 'right': [12, 12, 0, 12]},
    'info_three': {'a': [0, 8, 0, 12], 'b': [8, 8, 0, 12], 'c': [16, 8, 0, 12]},
    'info_four': {'a': [0, 6, 0, 12], 'b': [6, 6, 0, 12], 'c': [12, 6, 0, 12], 'd': [18, 6, 0, 12]},
    'info_quad': {'a': [0, 12, 0, 6], 'b': [12, 12, 0, 6], 'c': [0, 12, 6, 6], 'd': [12, 12, 6, 6]},
    'info_rows3': {'r1': [0, 24, 0, 4], 'r2': [0, 24, 4, 4], 'r3': [0, 24, 8, 4]},
    # 并列 + 底部通栏
    'info_split_strip': {'left': [0, 12, 0, 9], 'right': [12, 12, 0, 9], 'bottom': [0, 24, 9, 3]},
    'info_three_strip': {'a': [0, 8, 0, 9], 'b': [8, 8, 0, 9], 'c': [16, 8, 0, 9], 'bottom': [0, 24, 9, 3]},
    # 组合
    'info_mosaic': {'left': [0, 9, 0, 12], 'top': [9, 15, 0, 6], 'bottom': [9, 15, 6, 6]},
    'info_top_two': {'top': [0, 24, 0, 6], 'bl': [0, 12, 6, 6], 'br': [12, 12, 6, 6]},
    'info_two_top': {'tl': [0, 12, 0, 6], 'tr': [12, 12, 0, 6], 'bottom': [0, 24, 6, 6]},
    'info_stack': {'top': [0, 24, 0, 5], 'bottom': [0, 24, 5, 7]},
    'info_compare': {'left': [0, 11, 0, 12], 'mid': ([11, 2, 0, 12], 'none'), 'right': [13, 11, 0, 12]},
    'info_hub': {'lt': [0, 9, 0, 6], 'lb': [0, 9, 6, 6], 'center': ([9, 6, 0, 12], 'none'),
                 'rt': [15, 9, 0, 6], 'rb': [15, 9, 6, 6]},
    'info_grid': {'main': [0, 16, 0, 12], 'side': [16, 8, 0, 12]},
    'info_table_fig': {'left': [0, 13, 0, 12], 'right': [13, 11, 0, 12]},
}


def grid_box(cb, g, cols=24, rows=12, gx=0.14, gy=0.12):
    x, y, w, h = cb
    c, cs, r, rs = g
    cw = (w - gx * (cols - 1)) / cols
    rh = (h - gy * (rows - 1)) / rows
    return (x + c * (cw + gx), y + r * (rh + gy), cs * cw + (cs - 1) * gx, rs * rh + (rs - 1) * gy)


# ---------------------------------------------------------------- 基础绘制
def rect(slide, x, y, w, h, fill=None, line=None, lw=0.6, shape=MSO_SHAPE.RECTANGLE, dash=False, name=None):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        item.fill.background()
    if line:
        item.line.color.rgb = RGBColor.from_string(line)
        item.line.width = Pt(lw)
        if dash:
            item.line.dash_style = MSO_LINE.DASH
    else:
        item.line.fill.background()
    style = item._element.find(qn('p:style'))
    if style is not None:
        item._element.remove(style)
    if name:
        item.name = name
    return item


def _alpha(shape, value):
    clr = shape.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
    a = OxmlElement('a:alpha')
    a.set('val', str(int(value * 100000)))
    clr.append(a)


def runs_from(value, highlights=(), strong=(), color=INK, bold=False):
    red, black = list(highlights or []), list(strong or [])
    terms = sorted(red + black, key=len, reverse=True)
    parts = re.split('(' + '|'.join(re.escape(t) for t in terms) + ')', value) if terms else [value]
    out = []
    for part in parts:
        if not part:
            continue
        out.append((part, RED, True) if part in red else (part, INK, True) if part in black else (part, color, bold))
    return out


def rich(slide, paras, x, y, w, h, size, align=PP_ALIGN.LEFT, anchor=None, spacing=1.15, name=None):
    """paras: [{'runs': [(text, color, bold[, size])], 'size', 'indent', 'after', 'align', 'spacing'}]"""
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        item.name = name
    tf = item.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == 'middle' else MSO_ANCHOR.TOP
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get('align', align)
        p.line_spacing = para.get('spacing', spacing)
        p.space_after = Pt(para.get('after', 2))
        props = p._p.get_or_add_pPr()
        props.set('latinLnBrk', '0')
        props.set('eaLnBrk', '1')
        props.set('hangingPunct', '1')
        if para.get('indent'):
            props.set('marL', str(Inches(para['indent'])))
            props.set('indent', str(-Inches(para['indent'])))
        for run_spec in para['runs']:
            value, color, bold = run_spec[:3]
            run = p.add_run()
            run.text = value
            run.font.name = '微软雅黑' if value.strip() in ('■', '●', '✓') else 'Arial'   # 符号统一用中文字体，大小一致
            run.font.size = Pt(run_spec[3] if len(run_spec) > 3 else para.get('size', size))
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
            rpr = zh_run(run._r)
            ea = OxmlElement('a:ea')
            ea.set('typeface', '微软雅黑')
            rpr.append(ea)
    return item


def tx(slide, value, x, y, w, h, size, color=INK, bold=False, align='left', anchor='middle',
       highlights=None, spacing=1.12, name=None):
    paras = [{'runs': runs_from(line, highlights, (), color, bold)} for line in str(value).split('\n')]
    return rich(slide, paras, x, y, w, h, size, PP_ALIGN.CENTER if align == 'center' else PP_ALIGN.LEFT,
                anchor, spacing, name)


def fit_size(spec, kind, height_fn, w, h):
    """写了 size 就用；否则用默认字号（或 max_size），放不下才往下收到 12pt。"""
    if spec.get('size'):
        return spec['size']
    top = spec.get('max_size', SIZE[kind])
    s = top
    while s >= AUTOFIT_MIN:
        if height_fn(s, w) <= h + 0.02:
            return s
        s = round(s - 0.6, 1)
    return AUTOFIT_MIN


def _norm(spec):
    return spec if isinstance(spec, dict) else {'text': spec}


def _dark(color):
    r, g, b = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b < 150


def _arrow_line(slide, x1, y1, x2, y2, color=NAVY, width=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = OxmlElement('a:tailEnd')
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    ln.append(tail)
    c.name = '连接箭头'
    return c


def _circle_num(slide, n, x, y, d, fill=NAVY, color='FFFFFF'):
    rect(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL, name='编号圆')
    tx(slide, str(n), x, y, d, d, max(11, round(d * 72 * 0.55, 1)), color, True, 'center', name='编号')


# ---------------------------------------------------------------- 页面骨架
def headline(slide, title, meta):
    m = re.match(r'^\s*(\d+(?:\.\d+)*)\s+(.*)$', title)
    num, rest = (m.group(1), m.group(2)) if m else ('', title)
    x = 0.28
    if num:
        nw = text_width(num) * SIZE['num'] / 72 * 1.02 + 0.08
        tx(slide, num, x, 0.12, nw + 0.1, 0.62, SIZE['num'], NAVY, True, name='页标题编号')
        rect(slide, x + nw + 0.16, 0.17, 0.02, 0.5, NAVY, name='标题分隔线')
        x = x + nw + 0.36
    tx(slide, rest, x, 0.12, 10.1 - x, 0.62, SIZE['title'], NAVY, True, name='页标题')
    rect(slide, 0, 0.8, 13.333, 0.016, RULE, name='页眉分隔线')
    brand = ASSETS / 'branding' / 'xmu-energy-header.png'
    logo = meta.get('header_logo') or (str(brand) if meta.get('brand') == 'xmu' else None)
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(10.0), Inches(0.15), height=Inches(0.5))
        pic.left = Inches(13.02) - pic.width
        pic.name = '机构标识'
    elif meta.get('institution'):
        iw = text_width(meta['institution']) * 13.8 / 72 * 1.05 + 0.05
        rect(slide, 13.02 - iw - 0.42, 0.2, 0.33, 0.33, NAVY, name='单位标识')
        tx(slide, meta['institution'], 13.02 - iw, 0.2, iw + 0.1, 0.33, 13.8, NAVY, True, name='单位')


def lead(slide, spec, y):
    """✓ 引导句（深蓝细框）：一到两行，说清这页的背景与定位；tag 时左侧换成深蓝标签。"""
    spec = _norm(spec)
    size = spec.get('size', SIZE['lead'])
    tag_w = (text_width(spec['tag']) * size / 72 * 1.05 + 0.36) if spec.get('tag') else 0
    left = (tag_w + 0.18) if tag_w else 0.5
    tw = W0 - left - 0.2
    h = max(0.62, measure(spec['text'], size, tw, True, spacing=1.15)[1] + 0.16)
    rect(slide, X0, y, W0, h, 'FFFFFF', NAVY, name='引导句框')
    if tag_w:
        rect(slide, X0, y, tag_w, h, NAVY, name='引导句标签')
        tx(slide, spec['tag'], X0, y, tag_w, h, size, 'FFFFFF', True, 'center')
    else:
        tx(slide, '✓', X0 + 0.13, y, 0.36, h, SIZE['check'], NAVY, False, name='引导句标记')
    rich(slide, [{'runs': runs_from(spec['text'], spec.get('highlights'), spec.get('strong'), INK, True)}],
         X0 + left, y, tw, h, size, anchor='middle', name='引导句')
    return y + h


def sublead(slide, spec, y):
    """无框、居中、深蓝加粗的副引导句（参考图 2.1 / 2.2 / 3.4）：一句话点出研究路径。"""
    spec = _norm(spec)
    size = spec.get('size', SIZE['sublead'])
    h = measure(spec['text'], size, W0 - 0.4, True)[1] + 0.06
    rich(slide, [{'runs': runs_from(spec['text'], spec.get('highlights'), spec.get('strong'), NAVY, True)}],
         X0 + 0.2, y, W0 - 0.4, h, size, PP_ALIGN.CENTER, 'middle', name='副引导句')
    return y + h


def band(slide, value, x, y, w, h=BAND_H, size=None, num=None):
    rect(slide, x, y, w, h, NAVY, name='标题带')
    if num is not None:
        d = h * 0.72
        _circle_num(slide, num, x + 0.12, y + (h - d) / 2, d, 'FFFFFF', NAVY)
    tx(slide, value, x + 0.1 + (h if num is not None else 0), y, w - 0.2 - (h if num is not None else 0), h,
       size or SIZE['band'], 'FFFFFF', True, 'center', name='标题带文字')


def image_band(slide, spec, x, y, w, h):
    """图片衬底的标题带（参考图 4.2）：images 横向平铺并裁切，盖半透明深蓝 + 白粗字。"""
    files = [str(_asset(f)) for f in spec.get('images', []) if Path(_asset(f)).exists()]
    n = max(len(files), 1)
    cw = w / n
    for i, f in enumerate(files):
        with Image.open(f) as im:
            iw, ih = im.size
        pic = slide.shapes.add_picture(f, Inches(x + i * cw), Inches(y), Inches(cw), Inches(h))
        target, ratio = cw / h, iw / ih
        if ratio > target:
            pic.crop_left = pic.crop_right = (1 - target / ratio) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - ratio / target) / 2
        pic.name = '标题带底图'
    cover = rect(slide, x, y, w, h, NAVY, name='标题带')
    if files:
        _alpha(cover, spec.get('alpha', 0.8))
    tx(slide, spec['text'], x + 0.15, y, w - 0.3, h, spec.get('size', SIZE['band']), 'FFFFFF', True, 'center',
       name='标题带文字')


def summary(slide, spec, y, h=SUMMARY_H):
    """■ 底部结论条：本页要带走的一句话，红字标落脚词。"""
    spec = _norm(spec)
    size = spec.get('size', SIZE['summary'])
    rect(slide, X0, y, W0, h, 'FFFFFF', LINE, name='总结框')
    tx(slide, '■', X0 + 0.15, y, 0.3, h, size, INK, True, name='总结标记')
    rich(slide, [{'runs': runs_from(spec['text'], spec.get('highlights'), spec.get('strong'), INK, True)}],
         X0 + 0.5, y, W0 - 0.65, h, size, anchor='middle', name='总结条')


# ---------------------------------------------------------------- 文字组件
def _para_h(spec, size, w):
    bold = spec.get('bold', spec.get('color') in ('navy', 'red', 'white'))
    return sum(measure(l, size, w, bold, spacing=1.15)[1] for l in spec['text'].split('\n') if l) + 0.04


def para(slide, spec):
    """无框文字：说明句、红色小标题、深蓝小标题。color ink/navy/red/white，align center（默认）/left。"""
    x, y, w, h = spec['box']
    _check_sentence(spec['text'], '说明文字')
    color = {'navy': NAVY, 'red': RED, 'white': 'FFFFFF', 'grey': GREY}.get(spec.get('color'), INK)
    bold = spec.get('bold', color not in (INK, GREY))
    size = fit_size(spec, 'para', lambda s, ww: _para_h(spec, s, ww), w, h)
    paras = [{'runs': runs_from(l, spec.get('highlights'), spec.get('strong'), color, bold)}
             for l in spec['text'].split('\n') if l]
    align = PP_ALIGN.LEFT if spec.get('align') == 'left' else PP_ALIGN.CENTER
    rich(slide, paras, x, y, w, h, size, align, 'middle', 1.15, name='说明文字')


def _note_h(spec, size, w):
    h = 0.22
    if spec.get('title'):
        h += measure(spec['title'], size + 0.6, w - 0.3, True, spacing=1.15)[1]
    if spec.get('text'):
        h += measure(spec['text'], size, w - 0.3, spec.get('bold', False), spacing=1.15)[1]
    return h


def note(slide, spec):
    """浅蓝说明块：title（深蓝加粗）+ text（红字强调），默认居中。
    tone: red（浅红底红框红粗字）/ plain（无底）；dashed: 虚线框；line: 边框色。"""
    x, y, w, h = spec['box']
    tone = spec.get('tone')
    if tone == 'red':
        rect(slide, x, y, w, h, spec.get('fill', 'FDEDED'), RED, 0.9, name='说明块')
    elif tone != 'plain':
        rect(slide, x, y, w, h, spec.get('fill', LIGHT), spec.get('line'), dash=spec.get('dashed', False), name='说明块')
    _check_sentence(spec.get('text', ''), '说明块')
    size = fit_size(spec, 'note', lambda s, ww: _note_h(spec, s, ww), w, h)
    align = PP_ALIGN.LEFT if spec.get('align') == 'left' else PP_ALIGN.CENTER
    color, bold = (RED, True) if tone == 'red' else (INK, spec.get('bold', False))
    paras = []
    if spec.get('title'):
        paras.append({'runs': [(spec['title'], RED if spec.get('title_color') == 'red' else NAVY, True)],
                      'size': size + 0.6, 'after': 3})
    for line in (spec.get('text') or '').split('\n'):
        if line:
            paras.append({'runs': runs_from(line, spec.get('highlights'), spec.get('strong'), color, bold)})
    rich(slide, paras, x + 0.15, y + 0.04, w - 0.3, h - 0.08, size, align, 'middle', 1.15, name='说明框')


def _bullet_items(items):
    return [it if isinstance(it, dict) else {'text': it} for it in items]


def _markers_paras(items, size, after, label_color=NAVY):
    paras = []
    for it in items:
        runs = [('■', NAVY, False, round(size * 0.6, 1)), ('  ', INK, False)]
        if it.get('label'):
            runs.append((it['label'] + '：', label_color, True))
        runs += runs_from(it.get('text', ''), it.get('highlights'), it.get('strong'))
        paras.append({'runs': runs, 'indent': size / 72 * 1.05, 'after': after})
    return paras


def _markers_h(items, size, w, after):
    total = 0.06
    for it in items:
        value = (it.get('label', '') + '：' if it.get('label') else '') + it.get('text', '')
        total += measure(value, size, w - size / 72 * 1.05, False, spacing=1.15)[1] + after / 72
    return total


def markers(slide, spec):
    """■ 方块要点：items ["…"] 或 [{label, text, highlights}] → “■ 研究基础：数据与方法支撑”。
    fill: "light"（或 note）时放进浅蓝块（模板 1.1 右下）。"""
    x, y, w, h = spec['box']
    items = _bullet_items(spec['items'])
    pad = 0.0
    if spec.get('fill'):
        rect(slide, x, y, w, h, LIGHT if spec['fill'] in ('light', 'note') else spec['fill'], name='要点底')
        pad = 0.16
    after = spec.get('after', 5)
    size = fit_size(spec, 'markers', lambda s, ww: _markers_h(items, s, ww, after), w - 2 * pad, h - 2 * pad)
    rich(slide, _markers_paras(items, size, after), x + pad, y + pad, w - 2 * pad, h - 2 * pad, size,
         anchor='middle', spacing=1.15, name='方块要点')


def _head(slide, value, x, y, w, h, num=None, style='fill', size=None, align='left'):
    """浅蓝小标题条（模板“方法思路 / 优点 / 局限 / 主要研究任务”）；style bar 时左侧加深蓝竖条。"""
    rect(slide, x, y, w, h, LIGHT, name='小标题条')
    off = 0.14
    if style == 'bar':
        rect(slide, x, y, 0.06, h, NAVY, name='小标题竖条')
        off = 0.2
    if num is not None:
        d = h * 0.86
        _circle_num(slide, num, x + 0.1, y + (h - d) / 2, d)
        off = 0.2 + d
    tx(slide, value, x + off, y, w - off - 0.1, h, size or SIZE['head'], NAVY, True, align, name='小标题')


def _section_h(spec, size, w):
    items = _bullet_items(spec.get('items', []))
    h = 0.4 + 0.08 if spec.get('head') else 0
    if items:
        h += _markers_h(items, size, w - 0.1, spec.get('after', 4))
    if spec.get('text'):
        h += measure(spec['text'], size, w - 0.2, False, spacing=1.15)[1] + 0.06
    return h


def section(slide, spec):
    """小标题条 + 方块要点（模板最常用的单元）：{head, items, text, num, head_style}。"""
    x, y, w, h = spec['box']
    size = fit_size(spec, 'markers', lambda s, ww: _section_h(spec, s, ww), w, h)
    cy = y
    if spec.get('head'):
        _head(slide, spec['head'], x, y, w, 0.4, spec.get('num'), spec.get('head_style', 'fill'), align=spec.get('head_align', 'left'))
        cy += 0.48
    rest = y + h - cy
    items = _bullet_items(spec.get('items', []))
    if items:
        rich(slide, _markers_paras(items, size, spec.get('after', 4)), x + 0.1, cy, w - 0.1,
             rest if not spec.get('text') else rest * 0.7, size, anchor='middle' if spec.get('valign') == 'middle' else None,
             spacing=1.15, name='方块要点')
    elif spec.get('text'):
        _check_sentence(spec['text'], '小节说明')
        rich(slide, [{'runs': runs_from(spec['text'], spec.get('highlights'), spec.get('strong'))}],
             x + 0.1, cy, w - 0.2, rest, size, anchor=None, spacing=1.15, name='小节说明')


def _chev_h(spec, size, w):
    total = 0.06
    for it in spec['items']:
        total += measure('» ' + it['title'], size + 0.6, w, True, spacing=1.15)[1] + 2 / 72
        if it.get('text'):
            total += measure(it['text'], size, w, False, spacing=1.15)[1] + 8 / 72
    return total


def chevrons(slide, spec):
    """» 小标题 + 说明：items [{title, text, highlights}]。tone red（默认）/ navy。"""
    x, y, w, h = spec['box']
    size = fit_size(spec, 'chevrons', lambda s, ww: _chev_h(spec, s, ww), w, h)
    c = NAVY if spec.get('tone') == 'navy' else RED
    paras = []
    for it in spec['items']:
        paras.append({'runs': [('» ', c, True), (it['title'], c, True)], 'size': size + 0.6, 'after': 1})
        if it.get('text'):
            paras.append({'runs': runs_from(it['text'], it.get('highlights'), it.get('strong')), 'after': 8})
    rich(slide, paras, x, y, w, h, size, anchor='middle', name='箭头要点')


def _label_w(spec):
    size = spec.get('label_size', SIZE['label'])
    widest = max(text_width(l) for it in spec['items'] for l in it['label'].split('\n'))
    return spec.get('label_w', widest * size / 72 * 1.05 + 0.3)


def _label_rows_h(spec, w):
    size = spec.get('size', SIZE['note'])
    lw = _label_w(spec)
    return sum(max(measure(it.get('text', ''), size, w - lw - 0.35, False, spacing=1.15)[1] + 0.18,
                   len(it['label'].split('\n')) * spec.get('label_size', SIZE['label']) / 72 * 1.25 + 0.18)
               for it in spec['items']) + spec.get('gap', 0.1) * (len(spec['items']) - 1)


def label_rows(slide, spec):
    """标签说明行（参考图 2.1 右下“现存问题 | 基于…”）：items [{label, text, highlights}]，
    左侧红色粗体标签（label 可 \\n 竖排，label_color navy 改深蓝），右侧一句说明，每行一个虚线框。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.1)
    rh = (h - gap * (n - 1)) / n
    lw = _label_w(spec)
    lc = NAVY if spec.get('label_color') == 'navy' else RED
    size = fit_size(spec, 'note', lambda s, ww: _label_rows_h(dict(spec, size=s), ww), w, h)
    for i, it in enumerate(items):
        ry = y + i * (rh + gap)
        rect(slide, x, ry, w, rh, spec.get('fill', 'F7F9FD'), LINE, 0.75, dash=True, name='说明行')
        tx(slide, it['label'], x + 0.06, ry, lw - 0.08, rh, spec.get('label_size', SIZE['label']), lc, True,
           'center', spacing=1.0, name='说明行标签')
        rect(slide, x + lw + 0.02, ry + 0.1, 0.012, rh - 0.2, LINE, name='说明行分隔')
        rich(slide, [{'runs': runs_from(it.get('text', ''), it.get('highlights'), it.get('strong'))}],
             x + lw + 0.15, ry, w - lw - 0.25, rh, size, anchor='middle', name='说明行文字')


def links(slide, spec):
    """短语 —— 关键词（参考图 2.2 左下）：items [["突破关键技术瓶颈", "自主可控"], …]。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.12)
    rh = min(0.46, (h - gap * (n - 1)) / n)
    y0 = y + (h - (rh * n + gap * (n - 1))) / 2
    size = spec.get('size', SIZE['note'])
    lw, rw = w * spec.get('split', 0.52), w * 0.34
    for i, (left, right) in enumerate(items):
        ry = y0 + i * (rh + gap)
        rect(slide, x, ry, lw, rh, 'FFFFFF', NAVY, 0.9, MSO_SHAPE.ROUNDED_RECTANGLE, dash=True, name='短语框')
        tx(slide, left, x + 0.1, ry, lw - 0.2, rh, size, INK, True, 'center')
        rect(slide, x + lw, ry + rh / 2 - 0.012, w - lw - rw, 0.024, RED, name='连接线')
        rect(slide, x + w - rw, ry, rw, rh, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name='关键词')
        tx(slide, right, x + w - rw + 0.06, ry, rw - 0.12, rh, size, 'FFFFFF', True, 'center')


def callout(slide, spec):
    """一句关键论断：左侧红色竖条 + 加粗论断（默认 18pt），红字标关键词。"""
    x, y, w, h = spec['box']
    rect(slide, x, y, w, h, spec.get('fill', 'FFF5F5'), name='论断框')
    rect(slide, x, y, 0.08, h, RED, name='论断竖条')
    size = fit_size(dict(spec, max_size=spec.get('max_size', 18)), 'note',
                    lambda s, ww: measure(spec['text'], s, ww, True, spacing=1.2)[1] + 0.2, w - 0.45, h)
    rich(slide, [{'runs': runs_from(spec['text'], spec.get('highlights'), spec.get('strong'), INK, True)}],
         x + 0.28, y, w - 0.42, h, size, anchor='middle', spacing=1.2, name='论断')


def hexes(slide, spec):
    """一排六边形关键词（白底深蓝边，2–4 字，可 \\n 分两行）。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.1)
    hw = min((w - gap * (n - 1)) / n, spec.get('max_w', 1.3), h / 0.8)
    hh = hw * 0.8
    size = spec.get('size') or max(12, min(SIZE['hexes'], hw * 72 / 2.7))
    start = x + (w - (hw * n + gap * (n - 1))) / 2
    for i, label in enumerate(items):
        hx = start + i * (hw + gap)
        rect(slide, hx, y + (h - hh) / 2, hw, hh, 'FFFFFF', NAVY, 0.96, MSO_SHAPE.HEXAGON,
             name='六边形-' + label.replace('\n', ''))
        tx(slide, label, hx + 0.1, y + (h - hh) / 2, hw - 0.2, hh, size, INK, True, 'center', spacing=1.0)


def tags(slide, spec):
    """一排白底深蓝框标签（比六边形能放更长的词）：items ["…"]，可 \\n 分两行。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.12)
    tw = (w - gap * (n - 1)) / n
    th = min(h, spec.get('tag_h', 0.8))
    longest = max(text_width(l) for s in items for l in s.split('\n'))
    size = spec.get('size') or max(12, min(SIZE['note'], (tw - 0.2) * 72 / (longest * 1.05)))
    for i, label in enumerate(items):
        lx = x + i * (tw + gap)
        rect(slide, lx, y + (h - th) / 2, tw, th, 'FFFFFF', NAVY, 0.9, name='标签框')
        tx(slide, label, lx + 0.06, y + (h - th) / 2, tw - 0.12, th, size, NAVY, True, 'center', spacing=1.0)


def pill(slide, spec):
    """深蓝块白粗字（“关键问题识别与研究目标”）；full: true 撑满宽度，light: true 改浅蓝底深蓝字。"""
    x, y, w, h = spec['box']
    size = spec.get('size', SIZE['pill'])
    pw = w if spec.get('full') else min(w, text_width(spec['text']) * size / 72 * 1.06 + 0.6)
    ph = min(h, size / 72 * 1.9)
    px = x + (w - pw) / 2
    light = spec.get('light')
    rect(slide, px, y + (h - ph) / 2, pw, ph, LIGHT if light else NAVY, name='标签-' + spec['text'][:10])
    tx(slide, spec['text'], px + 0.08, y + (h - ph) / 2, pw - 0.16, ph, size, NAVY if light else 'FFFFFF', True,
       'center')


def arrow(slide, spec):
    """浅蓝大箭头：direction right / left / down / up，可带 label。"""
    x, y, w, h = spec['box']
    d = spec.get('direction', 'right')
    shape_type = {'right': MSO_SHAPE.RIGHT_ARROW, 'left': MSO_SHAPE.LEFT_ARROW,
                  'down': MSO_SHAPE.DOWN_ARROW, 'up': MSO_SHAPE.UP_ARROW}[d]
    s = spec.get('size', 0.55)
    aw, ah = (min(w, s), min(h, s * 0.8)) if d in ('right', 'left') else (min(w, s * 0.8), min(h, s))
    rect(slide, x + (w - aw) / 2, y + (h - ah) / 2, aw, ah, spec.get('fill', ARROW), shape=shape_type, name='大箭头')
    if spec.get('label'):
        tx(slide, spec['label'], x - 0.4, y + (h + ah) / 2 + 0.03, w + 0.8, 0.32, 13, NAVY, True, 'center')


def side_tag(slide, spec):
    """竖排深蓝标签（“现状 / 目标”），贴在一组说明的一侧。"""
    x, y, w, h = spec['box']
    tw = min(w, 0.44)
    rect(slide, x + (w - tw) / 2, y, tw, h, NAVY, name='竖标签')
    tx(slide, '\n'.join(spec['text']), x + (w - tw) / 2, y, tw, h, spec.get('size', SIZE['side']), 'FFFFFF',
       True, 'center', spacing=1.0)


def process(slide, spec):
    """流程条：items ["明确目标", "比较评估", "选择应用"]。style chevron（默认）/ pill（圆角块 + 小箭头）；
    colors 自定颜色（浅色自动配深蓝字），gradient: true 由浅到深（模板 2.1 右下）。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    hh = min(h, spec.get('bar_h', 0.44))
    y0 = y + (h - hh) / 2
    size = spec.get('size', SIZE['process'])
    colors = spec.get('colors') or (['DCE6F7', 'BCD3F6', '7FA3E0', '3E73CB', NAVY][-n:] if spec.get('gradient')
                                    else [NAVY] * n)
    if spec.get('style') == 'pill':
        aw = 0.34
        iw = (w - aw * (n - 1)) / n
        for i, label in enumerate(items):
            ix = x + i * (iw + aw)
            c = colors[i % len(colors)]
            rect(slide, ix, y0, iw, hh, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name='流程块')
            fit = min(size, max(11, (iw - 0.22) * 72 / (max(text_width(l) for l in label.split('\n')) * 1.06)))
            tx(slide, label, ix + 0.08, y0, iw - 0.16, hh, round(fit, 1), 'FFFFFF' if _dark(c) else NAVY, True, 'center', spacing=1.0)
            if i < n - 1:
                rect(slide, ix + iw + 0.07, y0 + hh * 0.22, aw - 0.14, hh * 0.56, NAVY,
                     shape=MSO_SHAPE.RIGHT_ARROW, name='流程箭头')
        return
    depth = hh * 0.34
    iw = (w + (n - 1) * (depth - 0.05)) / n
    for i, label in enumerate(items):
        ix = x + i * (iw - depth + 0.05)
        c = colors[i % len(colors)]
        shp = rect(slide, ix, y0, iw, hh, c, shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON,
                   name='流程箭头')
        try:
            shp.adjustments[0] = 0.34
        except Exception:
            pass
        left = ix + (depth if i else 0.08)
        bw = ix + iw - depth - left
        fit = min(size, max(11, (bw - 0.06) * 72 / (max(text_width(l) for l in label.split('\n')) * 1.06)))
        if '\n' not in label and fit < size * 0.8:
            print(f'  ! [{layouts.CURRENT_PAGE}] 流程条“{label}”太长，字号被压到 {fit:.0f}pt：缩短到 {int(bw * 72 / size / 1.06)} 字以内')
        tx(slide, label, left, y0, bw, hh, round(fit, 1), 'FFFFFF' if _dark(c) else NAVY, True, 'center', spacing=1.0)


# ---------------------------------------------------------------- 图
def _asset(f):
    s = str(f)
    return ASSETS / s[2:] if s.startswith('@/') else Path(s)


def _fit_rect(spec):
    x, y, w, h = spec['box']
    f = spec.get('file')
    if not f or not _asset(f).exists() or spec.get('fit') == 'cover':
        return x, y, w, h
    with Image.open(_asset(f)) as im:
        iw, ih = im.size
    cl, ct, cr, cb = spec.get('crop', [0, 0, 0, 0])
    ratio = iw * (1 - cl - cr) / (ih * (1 - ct - cb))
    fw = min(w, h * ratio)
    fh = fw / ratio
    return x + (w - fw) / 2, y + (h - fh) / 2, fw, fh


def figure(slide, spec, block):
    """info 页的图：默认无边框、按比例尽量放大。
    caption_style: box（图注放进浅蓝块贴底，可加 caption_title 两级图注）/ 默认（图下一行粗体图注）；
    overlay: {text, highlights} 在图下部压半透明深蓝说明（参考图 3.4）。"""
    spec = dict(spec)
    block = getattr(block, '_raw', block)
    spec.setdefault('caption', '')          # 不写 caption 就没有图注（不再把 label 当图注）
    if spec.get('file'):
        spec['file'] = str(_asset(spec['file']))
    spec.setdefault('frame', False)
    spec.setdefault('caption_size', SIZE['caption'])
    if spec.get('caption'):
        spec.setdefault('caption_bold', True)
    x, y, w, h = spec['box']
    _check_figure_size(spec)
    if spec.get('caption_style') == 'box' and (spec.get('caption') or spec.get('caption_title')):
        paras = []
        if spec.get('caption_title'):
            paras.append({'runs': [(spec['caption_title'], NAVY, True)], 'after': 2})
        if spec.get('caption'):
            paras.append({'runs': runs_from(spec['caption'], spec.get('caption_highlights'), (),
                                            NAVY if not spec.get('caption_title') else INK,
                                            not spec.get('caption_title'))})
        ch = 0.14 + sum(measure(''.join(r[0] for r in p['runs']), SIZE['caption'], w - 0.2, True)[1] for p in paras)
        ch = max(0.36, ch)
        block(slide, dict(spec, caption='', box=[x, y, w, h - ch - 0.06]))
        rect(slide, x, y + h - ch, w, ch, LIGHT, name='图注框')
        rich(slide, paras, x + 0.1, y + h - ch, w - 0.2, ch, SIZE['caption'], PP_ALIGN.CENTER, 'middle', name='图注')
        return
    if spec.get('overlay'):
        ov = _norm(spec['overlay'])
        block(slide, dict(spec, caption=''))
        fx, fy, fw, fh = _fit_rect(spec)
        size = ov.get('size', SIZE['note'])
        oh = measure(ov['text'], size, fw - 0.5, True, spacing=1.2)[1] + 0.2
        shade = rect(slide, fx + 0.08, fy + fh - oh - 0.08, fw - 0.16, oh, NAVY, name='图上说明底')
        _alpha(shade, 0.9)
        runs = [(t, 'FFE08A' if c == RED else 'FFFFFF', True)
                for t, c, _ in runs_from(ov['text'], ov.get('highlights'), (), 'FFFFFF', True)]
        rich(slide, [{'runs': runs}], fx + 0.25, fy + fh - oh - 0.08, fw - 0.5, oh, size, anchor='middle',
             spacing=1.2, name='图上说明')
        return
    block(slide, spec)


def _place_figure(slide, fig, rect_, block):
    fig = dict(fig)
    fig.setdefault('caption', '')
    fig['type'] = fig.get('type', 'figure')
    fig['box'] = list(rect_)
    if fig['caption'] and not fig.get('caption_style'):
        fig['box'][3] -= 0.58
    figure(slide, fig, block)


def fig_list(slide, spec, block):
    """图 + 要点竖排列表（参考图 3.1 右栏）：items [{figure, title, text, highlights}]。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.1)
    rh = (h - gap * (n - 1)) / n
    fw = w * spec.get('fig_w', 0.36)
    size = spec.get('size', SIZE['markers'])
    for i, it in enumerate(items):
        ry = y + i * (rh + gap)
        if it.get('figure'):
            _place_figure(slide, it['figure'], [x, ry, fw, rh], block)
        paras = [{'runs': [('●  ', NAVY, True, size * 0.7), (it['title'], NAVY, True)], 'size': size + 1, 'after': 2}]
        if it.get('text'):
            paras.append({'runs': runs_from(it['text'], it.get('highlights'), it.get('strong'))})
        rich(slide, paras, x + fw + 0.15, ry, w - fw - 0.15, rh, size, anchor='middle', name='图注要点')


def card(slide, spec, block):
    """图 + 标题 + 红点要点的卡片（参考图 3.4 负责人卡；也作平台卡、工具卡、样品卡）。"""
    x, y, w, h = spec['box']
    fw = w * spec.get('fig_w', 0.34)
    if spec.get('figure'):
        _place_figure(slide, dict(spec['figure'], fit=spec['figure'].get('fit', 'cover')), [x, y, fw, h], block)
    size = spec.get('size', SIZE['markers'])
    paras = [{'runs': [(spec['title'], NAVY, True)], 'size': size + 3, 'after': 5}]
    for it in spec.get('items', []):
        it = _norm(it)
        paras.append({'runs': [('●  ', RED if spec.get('items_red', True) else NAVY, True, size * 0.7)] +
                      runs_from(it['text'], it.get('highlights'), it.get('strong'),
                                RED if spec.get('items_red', True) else INK, True),
                      'indent': size / 72 * 1.0, 'after': 3})
    if spec.get('text'):
        paras.append({'runs': runs_from(spec['text'], spec.get('highlights'), spec.get('strong'))})
    rich(slide, paras, x + fw + 0.2, y, w - fw - 0.25, h, size, anchor='middle', name='卡片文字')


# ---------------------------------------------------------------- 结构组件
def stages(slide, spec, block):
    """横向流程 / 技术路线：items [{title, figure, items, text, output}]，阶段之间自动加浅蓝箭头。
    每个阶段：深蓝标题条（head_style light 时浅蓝）→ 图 → 要点 → 浅蓝“输出”条；frame: false 去掉阶段框。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    aw = spec.get('arrow_w', 0.34)
    sw = (w - (n - 1) * aw) / n
    size = spec.get('size', SIZE['markers'])
    light = spec.get('head_style') == 'light'
    bullets = [_bullet_items(it.get('items', [])) for it in items]
    bul_h = max([_markers_h(b, size, sw - 0.3, 3) for b in bullets if b] or [0])
    txt_h = max([measure(it['text'], size, sw - 0.3, False, spacing=1.15)[1] + 0.2 for it in items if it.get('text')] or [0])
    out_h = 0.4 if any(it.get('output') for it in items) else 0
    head_h = 0.44 if any(it.get('title') for it in items) else 0
    for i, it in enumerate(items):
        sx = x + i * (sw + aw)
        if spec.get('frame', True):
            rect(slide, sx, y + head_h, sw, h - head_h, 'FFFFFF', LINE, name='阶段框')
        if it.get('title'):
            rect(slide, sx, y, sw, head_h, LIGHT if light else NAVY, name='阶段标题')
            tx(slide, it['title'], sx + 0.06, y, sw - 0.12, head_h, SIZE['head'], NAVY if light else 'FFFFFF', True,
               'center')
        cy = y + head_h + 0.1
        bottom = y + h - 0.1 - (out_h + 0.08 if out_h else 0) - (bul_h + 0.06 if bul_h else 0) - (txt_h if txt_h else 0)
        if it.get('figure'):
            _place_figure(slide, it['figure'], [sx + 0.1, cy, sw - 0.2, bottom - cy], block)
        by = bottom + 0.06
        if bullets[i]:
            rich(slide, _markers_paras(bullets[i], size, 3), sx + 0.15, by, sw - 0.25, bul_h, size, name='阶段要点')
            by += bul_h
        if it.get('text'):
            note(slide, {'box': [sx + 0.1, by, sw - 0.2, txt_h - 0.06], 'text': it['text'], 'size': size,
                         'highlights': it.get('highlights')})
        if it.get('output'):
            rect(slide, sx + 0.1, y + h - 0.1 - out_h, sw - 0.2, out_h, LIGHT, name='输出条')
            tx(slide, it['output'], sx + 0.12, y + h - 0.1 - out_h, sw - 0.24, out_h, SIZE['caption'] + 0.6, NAVY,
               True, 'center', highlights=it.get('output_highlights'))
        if i < n - 1:
            fy = cy + max(bottom - cy, 0.5) / 2
            arrow(slide, {'box': [sx + sw + 0.03, fy - 0.25, aw - 0.06, 0.5], 'direction': 'right', 'size': 0.3})


def timeline(slide, spec, block):
    """时间轴（模板 4.1 + 生图稿第三组）：items [{stage, title, time, blocks: [{head, items}], text, figure}]。
    上方浅蓝阶段框 → 带编号节点的轴线（末端箭头）→ 时间 → 阶段目标 / 预期成果等小节；阶段之间虚线分隔。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    seg = w / n
    size = spec.get('size', SIZE['markers'])
    has_head = any(it.get('stage') or it.get('title') for it in items)
    head_h = 0.66 if any(it.get('stage') and it.get('title') for it in items) else 0.4 if has_head else 0
    line_y = y + head_h + 0.32
    rect(slide, x, line_y - 0.025, w - 0.05, 0.05, NAVY, name='时间轴线')
    rect(slide, x + w - 0.2, line_y - 0.1, 0.2, 0.2, NAVY, shape=MSO_SHAPE.ISOSCELES_TRIANGLE, name='时间轴箭头').rotation = 90
    for i, it in enumerate(items):
        cx = x + seg * (i + 0.5)
        lx, lw = cx - seg / 2 + 0.08, seg - 0.16
        if i:
            sep = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + seg * i), Inches(y),
                                             Inches(x + seg * i), Inches(y + h))
            sep.line.color.rgb = RGBColor.from_string('B8C7DE')
            sep.line.width = Pt(0.9)
            sep.line.dash_style = MSO_LINE.DASH
            sep.name = '阶段分隔'
        if has_head:
            rect(slide, lx, y, lw, head_h, LIGHT, name='阶段框')
            paras = []
            if it.get('stage'):
                paras.append({'runs': [(it['stage'], NAVY, True)], 'size': SIZE['head'] - 1})
            if it.get('title'):
                paras.append({'runs': [(it['title'], NAVY, True)]})
            rich(slide, paras, lx + 0.05, y, lw - 0.1, head_h, SIZE['head'], PP_ALIGN.CENTER, 'middle')
        d = 0.34
        rect(slide, cx - d / 2, line_y - d / 2, d, d, 'FFFFFF', NAVY, 2.5, MSO_SHAPE.OVAL, name='时间节点')
        rect(slide, cx - 0.07, line_y - 0.07, 0.14, 0.14, NAVY, shape=MSO_SHAPE.OVAL, name='时间节点心')
        cy = line_y + 0.24
        if it.get('time'):
            tx(slide, it['time'], lx, cy, lw, 0.32, 13.2, GREY, False, 'center', name='时间')
            cy += 0.38
        rest = y + h - cy
        blocks = it.get('blocks', [])
        if it.get('figure'):
            fh = rest * spec.get('fig_share', 0.45) if (blocks or it.get('text')) else rest
            _place_figure(slide, it['figure'], [lx, cy, lw, fh], block)
            cy += fh + 0.08
            rest = y + h - cy
        if blocks:
            bh = (rest - 0.08 * (len(blocks) - 1)) / len(blocks)
            for b in blocks:
                section(slide, dict(b, box=[lx, cy, lw, bh], head_style=b.get('head_style', 'fill'),
                                    size=b.get('size', spec.get('size'))))
                cy += bh + 0.08
        elif it.get('text'):
            note(slide, {'box': [lx, cy, lw, rest], 'title': it.get('note_title'), 'text': it['text'], 'size': size,
                         'highlights': it.get('highlights')})


def rows(slide, spec, block):
    """编号结论行：items [{title, text, highlights, figure}]，左编号、中文字、右侧可带小图。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.12)
    rh = (h - gap * (n - 1)) / n
    size = spec.get('size', SIZE['note'])
    for i, it in enumerate(items):
        ry = y + i * (rh + gap)
        rect(slide, x, ry, w, rh, LIGHT, name='结论行')
        d = min(0.5, rh - 0.2)
        _circle_num(slide, i + 1, x + 0.18, ry + (rh - d) / 2, d)
        fw = min(w * 0.34, (rh - 0.12) * 1.78) if it.get('figure') else 0
        tx0 = x + 0.18 + d + 0.25
        tw = w - (tx0 - x) - 0.2 - (fw + 0.2 if fw else 0)
        paras = [{'runs': [(it['title'], NAVY, True)], 'size': size + 1.5, 'after': 3}]
        for t in _bullet_items(it.get('items', [])):          # 结论下面可以再分点
            paras.append({'runs': [('■', NAVY, False, round(size * 0.6, 1)), ('  ', INK, False)] +
                          runs_from(t['text'], t.get('highlights'), t.get('strong')), 'indent': size / 72 * 1.05, 'after': 2})
        if it.get('text'):
            _check_sentence(it['text'], '结论行')
            paras.append({'runs': runs_from(it['text'], it.get('highlights'), it.get('strong'))})
        rich(slide, paras, tx0, ry + 0.05, tw, rh - 0.1, size, anchor='middle', name='结论文字')
        if fw:
            _place_figure(slide, it['figure'], [x + w - fw - 0.12, ry + 0.06, fw, rh - 0.12], block)


def layers(slide, spec):
    """层级条（体系、层次推进）：items [{title, text}]，自上而下由深到浅、由宽到窄（pyramid: false 时等宽）。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    gap = spec.get('gap', 0.1)
    lh = (h - gap * (n - 1)) / n
    size = spec.get('size', SIZE['note'])
    for i, it in enumerate(items):
        shrink = (0.18 * i / max(n - 1, 1)) if spec.get('pyramid', True) else 0
        lw = w * (1 - shrink)
        lx = x + (w - lw) / 2
        rect(slide, lx, y + i * (lh + gap), lw, lh, LAYER[min(i, len(LAYER) - 1)], name='层级条')
        runs = [(it['title'], 'FFFFFF', True, size + 1.5)]
        if it.get('text'):
            runs.append(('　' + it['text'], 'FFFFFF', False))
        rich(slide, [{'runs': runs}], lx + 0.2, y + i * (lh + gap), lw - 0.4, lh, size, PP_ALIGN.CENTER, 'middle',
             name='层级文字')


def steps_v(slide, spec):
    """竖向编号步骤：圆形编号 + “标题：一句说明”，编号之间连竖线。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    sh = h / n
    d = min(0.44, sh * 0.7)
    size = spec.get('size', SIZE['note'])
    rect(slide, x + d / 2 - 0.015, y + sh / 2, 0.03, sh * (n - 1), NAVY, name='步骤连线')
    for i, it in enumerate(items):
        cy = y + i * sh
        _circle_num(slide, i + 1, x, cy + (sh - d) / 2, d)
        runs = [(it['title'], NAVY, True)]
        if it.get('text'):
            runs += [('：', NAVY, True)] + runs_from(it['text'], it.get('highlights'), it.get('strong'))
        rich(slide, [{'runs': runs}], x + d + 0.2, cy, w - d - 0.2, sh, size, anchor='middle', name='步骤文字')


def quadrant(slide, spec):
    """2×2 象限：items 4 个 {title, text}（左上、右上、左下、右下），可写 x_label / y_label。"""
    x, y, w, h = spec['box']
    lab = 0.4 if (spec.get('x_label') or spec.get('y_label')) else 0
    gx, gw, gh = x + lab, w - lab, h - lab
    gap = 0.1
    cw, ch = (gw - gap) / 2, (gh - gap) / 2
    fills = spec.get('fills', [LIGHT, 'D6E2F7', 'F3F6FB', LIGHT])
    for i, it in enumerate(spec['items'][:4]):
        note(slide, {'box': [gx + (i % 2) * (cw + gap), y + (i // 2) * (ch + gap), cw, ch], 'title': it.get('title'),
                     'text': it.get('text', ''), 'fill': fills[i], 'highlights': it.get('highlights')})
    if spec.get('x_label'):
        tx(slide, spec['x_label'] + '  →', gx, y + h - lab + 0.04, gw, 0.34, 14, NAVY, True, 'center')
    if spec.get('y_label'):
        t = tx(slide, spec['y_label'] + '  →', x - gh / 2 + 0.2, y + gh / 2 - 0.17, gh, 0.34, 14, NAVY, True, 'center')
        t.rotation = -90
        t.name = '标注-纵轴'


def _cell(slide, cell, x, y, w, h, block, label_style):
    """matrix 的单元格：文字 / {title, text} / {figure, caption} / {figure, items}。"""
    cell = _norm(cell)
    if cell.get('body'):                    # 一句说明：浅蓝底、黑字、左对齐（模板“研究动机”）
        rect(slide, x, y, w, h, LIGHT, name='单元格')
        size = fit_size(cell, 'note', lambda s, ww: measure(cell['body'], s, ww, False, spacing=1.15)[1] + 0.1, w - 0.4, h)
        rich(slide, [{'runs': runs_from(cell['body'], cell.get('highlights'), cell.get('strong'))}],
             x + 0.2, y, w - 0.4, h, size, anchor='middle', name='单元格文字')
        return
    if cell.get('figure'):
        if cell.get('items'):
            fw = w * 0.5
            _place_figure(slide, cell['figure'], [x, y, fw, h], block)
            items = _bullet_items(cell['items'])
            size = fit_size(cell, 'markers', lambda s, ww: _markers_h(items, s, ww, 2), w - fw - 0.1, h)
            rich(slide, _markers_paras(items, size, 2), x + fw + 0.1, y, w - fw - 0.1, h, size, anchor='middle')
        else:
            cap = 0.3 if cell.get('caption') else 0
            _place_figure(slide, dict(cell['figure'], caption=''), [x, y, w, h - cap], block)
            if cap:
                tx(slide, cell['caption'], x, y + h - cap, w, cap, SIZE['caption'], INK, True, 'center')
        return
    rect(slide, x, y, w, h, LIGHT, LINE, name='单元格')
    paras = [{'runs': [(cell.get('title', cell.get('text', '')), NAVY, True)], 'after': 1}]
    if cell.get('title') and cell.get('text'):
        paras.append({'runs': [(cell['text'], '3A4A63', False)], 'size': SIZE['caption']})
    size = cell.get('size') or min(SIZE['cell'], max(12, (w - 0.2) * 72 / max(text_width(paras[0]['runs'][0][0]), 1) / 1.05))
    rich(slide, paras, x + 0.08, y, w - 0.16, h, size, PP_ALIGN.CENTER, 'middle', name='单元格文字')


def matrix(slide, spec, block):
    """分层矩阵（总体技术路线 / 分层架构 / 多维对比）：
    rows [{label, sub, cells: [...], notes: [...]}]；cells 每格写文字、{title, text}、{figure, caption} 或 {figure, items}；
    columns: 表头（第一个是行标签列的表头）；arrows: down / up（行间细箭头）；top: 顶部目标条；bottom: 底部输出条；
    side: {title, items} 右侧支撑列；label_style navy（默认）/ light；row_frame: true 给每行加细框。"""
    x, y, w, h = spec['box']
    rows_ = spec['rows']
    n = len(rows_)
    ncell = max(len(r['cells']) for r in rows_)
    side_w = spec.get('side_w', 1.75) if spec.get('side') else 0
    notes_w = spec.get('notes_w', 2.6) if any(r.get('notes') for r in rows_) else 0
    if spec.get('row_frame'):                        # 每行加细框时内容往里收，框与页面其他元素对齐
        x, w = x + 0.08, w - 0.16
    lab_w = spec.get('label_w', 1.75)
    gx = spec.get('gap_x', 0.25)
    arrow_gap = 0.34 if spec.get('arrows') else 0.14
    top_h = 0.46 if spec.get('top') else 0
    bot_h = 0.4 if spec.get('bottom') else 0
    head_h = 0.4 if spec.get('columns') else 0
    body_x = x + lab_w + gx
    body_w = w - lab_w - gx - (notes_w + 0.15 if notes_w else 0) - (side_w + 0.3 if side_w else 0)
    cw = (body_w - (ncell - 1) * gx) / ncell
    cy = y
    if top_h:
        tw = body_w * 0.6
        rect(slide, body_x + (body_w - tw) / 2, cy, tw, top_h, NAVY, name='顶部目标')
        tx(slide, spec['top'], body_x + (body_w - tw) / 2, cy, tw, top_h, SIZE['head'] + 1, 'FFFFFF', True, 'center')
        cy += top_h
        _arrow_line(slide, body_x + body_w / 2, cy, body_x + body_w / 2, cy + 0.26)
        cy += 0.3
    if head_h:
        cols = spec['columns']
        rect(slide, x, cy, lab_w, head_h, NAVY, name='表头')
        tx(slide, cols[0], x, cy, lab_w, head_h, SIZE['head'], 'FFFFFF', True, 'center')
        for j, c in enumerate(cols[1:]):
            rect(slide, body_x + j * (cw + gx), cy, cw, head_h, NAVY, name='表头')
            tx(slide, c, body_x + j * (cw + gx), cy, cw, head_h, SIZE['head'], 'FFFFFF', True, 'center')
        cy += head_h + 0.1
    rows_top = cy
    rh = (y + h - (bot_h + 0.14 if bot_h else 0) - cy - arrow_gap * (n - 1)) / n
    for i, r in enumerate(rows_):
        ry = cy + i * (rh + arrow_gap)
        if spec.get('row_frame'):
            rect(slide, x - 0.08, ry - 0.05, w - (side_w + 0.3 if side_w else 0) + 0.16, rh + 0.1, None, LINE,
                 name='行框')
        navy = spec.get('label_style', 'navy') == 'navy'
        rect(slide, x, ry, lab_w, rh, NAVY if navy else LIGHT, None if navy else LINE, name='行标签')
        paras = [{'runs': [(r['label'], 'FFFFFF' if navy else NAVY, True)], 'size': SIZE['cell'] + 1}]
        if r.get('sub'):
            paras.append({'runs': [(r['sub'], 'FFFFFF' if navy else NAVY, True)], 'size': SIZE['caption']})
        rich(slide, paras, x + 0.05, ry, lab_w - 0.1, rh, SIZE['cell'], PP_ALIGN.CENTER, 'middle', name='行标签文字')
        for j, cell in enumerate(r['cells']):
            _cell(slide, cell, body_x + j * (cw + gx), ry, cw, rh, block, spec.get('label_style'))
            if spec.get('arrows') and i < n - 1:
                ax = body_x + j * (cw + gx) + cw / 2
                if spec['arrows'] == 'up':
                    _arrow_line(slide, ax, ry + rh + arrow_gap - 0.04, ax, ry + rh + 0.04)
                else:
                    _arrow_line(slide, ax, ry + rh + 0.04, ax, ry + rh + arrow_gap - 0.04)
        if r.get('notes'):
            items = _bullet_items(r['notes'])
            nx = x + w - (side_w + 0.3 if side_w else 0) - notes_w
            size = fit_size({}, 'markers', lambda s, ww: _markers_h(items, s, ww, 2), notes_w, rh)
            rich(slide, _markers_paras(items, size, 2), nx, ry, notes_w, rh, size, anchor='middle', name='行要点')
    if bot_h:
        by = y + h - bot_h
        rect(slide, body_x, by, body_w, bot_h, NAVY, name='底部输出')
        tx(slide, spec['bottom'], body_x, by, body_w, bot_h, SIZE['head'] - 0.6, 'FFFFFF', True, 'center')
    if side_w:
        sx = x + w - side_w
        sy, sh = rows_top, y + h - rows_top - (bot_h + 0.14 if bot_h else 0)
        side = spec['side']
        rect(slide, sx, sy, side_w, sh, 'FFFFFF', NAVY, 1.2, name='支撑列')
        rect(slide, sx, sy, side_w, 0.46, NAVY, name='支撑列标题')
        tx(slide, side['title'], sx, sy, side_w, 0.46, SIZE['head'], 'FFFFFF', True, 'center')
        its = side['items']
        ih = (sh - 0.46 - 0.1) / len(its)
        for k, it in enumerate(its):
            it = _norm(it)
            iy = sy + 0.5 + k * ih
            if it.get('sub'):
                rect(slide, sx + 0.1, iy + 0.05, side_w - 0.2, ih - 0.1, LIGHT, name='支撑项')
            paras = [{'runs': [(it['text'], NAVY if it.get('sub') else INK, bool(it.get('sub')))]}]
            if it.get('sub'):
                paras.append({'runs': [(it['sub'], '3A4A63', False)], 'size': SIZE['small']})
            rich(slide, paras, sx + 0.12, iy, side_w - 0.24, ih, SIZE['markers'], PP_ALIGN.CENTER, 'middle')


def cycle(slide, spec):
    """闭环（模板 3.1 中心）：items ["数据获取", "模型构建", "应用部署", "评估反馈"]（从左上起顺时针），center 中心文字。"""
    x, y, w, h = spec['box']
    d = min(w, h)
    cx, cy = x + w / 2, y + h / 2
    n = len(spec['items'])
    span = 360 / n
    colors = spec.get('colors') or ['0A3A8C', '1F5FC0', '3E7FD6', '8DB3EA'][:n]
    for i, label in enumerate(spec['items']):
        mid = (spec.get('start', 225) + i * span) % 360
        a1, a2 = (mid - span / 2 + 4) % 360, (mid + span / 2 - 4) % 360
        arc = rect(slide, cx - d / 2, cy - d / 2, d, d, colors[i % len(colors)], shape=MSO_SHAPE.BLOCK_ARC, name='闭环弧')
        arc.adjustments[0] = a1 * 0.6
        arc.adjustments[1] = a2 * 0.6
        arc.adjustments[2] = 0.22
        r = d * 0.39
        lx = cx + r * math.cos(math.radians(mid))
        ly = cy + r * math.sin(math.radians(mid))
        tx(slide, label, lx - 0.6, ly - 0.3, 1.2, 0.6, spec.get('size', 15), 'FFFFFF', True, 'center',
           name='闭环标签')
    inner = d * 0.5
    rect(slide, cx - inner / 2, cy - inner / 2, inner, inner, LIGHT, shape=MSO_SHAPE.OVAL, name='闭环中心')
    tx(slide, spec.get('center', ''), cx - inner / 2 + 0.06, cy - inner / 2, inner - 0.12, inner,
       spec.get('center_size', 17), NAVY, True, 'center', spacing=1.1)


# ---------------------------------------------------------------- 表格
def _cell_border(cell, sides, color, width_pt):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        if tag[-1] in sides:
            old = tcPr.find(qn(tag))
            if old is not None:
                tcPr.remove(old)
    for tag in ('a:lnB', 'a:lnT', 'a:lnR', 'a:lnL'):       # 倒序 insert(0) → 顺序 L R T B
        if tag[-1] not in sides:
            continue
        ln = OxmlElement(tag)
        ln.set('w', str(int(width_pt * 12700)))
        ln.set('cap', 'flat')
        ln.set('cmpd', 'sng')
        ln.set('algn', 'ctr')
        fill = OxmlElement('a:solidFill')
        clr = OxmlElement('a:srgbClr')
        clr.set('val', color)
        fill.append(clr)
        ln.append(fill)
        dash = OxmlElement('a:prstDash')
        dash.set('val', 'solid')
        ln.append(dash)
        tcPr.insert(0, ln)


def info_table(slide, spec):
    """info 页表格（模板 2.4 / 参考图 4.2）：表头深蓝白粗字，首列浅蓝底深蓝粗体，全部居中，细蓝灰框线；
    highlights 让单元格关键词变红，highlight_rows（从 1 数的正文行号）给整行加红框。"""
    x, y, w, h = spec['box']
    rows_ = [spec['columns']] + spec['rows']
    size = spec.get('size', SIZE['table'])
    shape = slide.shapes.add_table(len(rows_), len(rows_[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = '表格'
    table = shape.table
    tblPr = table._tbl.tblPr
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'   # No Style, Table Grid
    if spec.get('widths'):
        for column, weight in zip(table.columns, spec['widths']):
            column.width = Inches(w * weight / sum(spec['widths']))
    head_h = min(0.44, h / len(rows_))
    table.rows[0].height = Inches(head_h)
    for i in range(1, len(rows_)):
        table.rows[i].height = Inches((h - head_h) / (len(rows_) - 1))
    aligns = spec.get('align', ['center'] * len(rows_[0]))
    marked = set(spec.get('highlight_rows', []))
    if 0 in marked or any(r > len(spec['rows']) for r in marked):
        print(f"  ! [{layouts.CURRENT_PAGE}] highlight_rows 从 1 数正文行（不含表头），{sorted(marked)} 超出范围，已忽略越界的行号")
        marked = {r for r in marked if 1 <= r <= len(spec['rows'])}
    for i, row in enumerate(rows_):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            first = j == 0 and spec.get('first_col', True)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if i == 0 else LIGHT if first else 'FFFFFF')
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if aligns[j] == 'center' else PP_ALIGN.LEFT
            if i == 0:
                runs = [(str(value), 'FFFFFF', True)]
            elif first:
                runs = [(str(value), NAVY, True)]
            else:
                runs = runs_from(str(value), spec.get('highlights'), spec.get('strong'))
            for value_, color, bold in runs:
                run = p.add_run()
                run.text = value_
                run.font.name = 'Arial'
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = RGBColor.from_string(color)
                ea = OxmlElement('a:ea')
                ea.set('typeface', '微软雅黑')
                zh_run(run._r).append(ea)
            _cell_border(cell, 'LRTB', LINE if i else NAVY, 0.6)
            if i in marked:
                _cell_border(cell, 'TB' + ('L' if j == 0 else '') + ('R' if j == len(row) - 1 else ''), RED, 2.0)
            elif i + 1 in marked:
                _cell_border(cell, 'B', RED, 2.0)


# ---------------------------------------------------------------- 两条全局检查
def _check_sentence(text, where):
    """一句话说不清的内容要分点：单段超过 36 字且有 3 个以上分句时提示改成 markers / section。"""
    for para_ in str(text).split('\n'):
        clauses = [c for c in re.split('[，；。,;]', para_) if c.strip()]
        if len(para_) > 36 and len(clauses) >= 3:
            print(f'  ! [{layouts.CURRENT_PAGE} 分区 {layouts.CURRENT_REGION}] {where}“{para_[:16]}…”一句话说不清'
                  f'（{len(para_)} 字、{len(clauses)} 个分句）：拆成分点（markers / section）')


BIG_MIN = (7.5, 4.0)       # 大图最小宽 × 高（英寸）
_PAGE_BIG = [0]            # 本页大图计数


def _check_figure_size(spec):
    """role: "big"（框架图、技术路线、核心结果、系统架构、平台主界面）必须放得够大；
    role: "small"（辅助证据、示意、图标）不承载需要读的细节。"""
    role = spec.get('role')
    if role == 'small':
        fx, fy, fw, fh = _fit_rect(spec)
        if fw > 4.6:
            print(f'  · [{layouts.CURRENT_PAGE}] 小图“{spec.get("label", "")}”放到了 {fw:.1f} in 宽：'
                  f'小图只放图标、缩略、简单示意；内容需要细看就改成中图（去掉 role: small）')
        return
    if role not in ('big', 'main'):
        return
    _PAGE_BIG[0] += 1
    fx, fy, fw, fh = _fit_rect(spec)
    if fw < BIG_MIN[0] - 0.05 and fh < BIG_MIN[1] - 0.05:
        print(f'  ! [{layouts.CURRENT_PAGE} 分区 {layouts.CURRENT_REGION}] 大图“{spec.get("label", "")}”只有 '
              f'{fw:.1f}×{fh:.1f} in（要求宽 ≥ {BIG_MIN[0]} 或高 ≥ {BIG_MIN[1]}）：改用 info_big / info_hero，'
              f'去掉标题带，或删掉同栏文字')


# ---------------------------------------------------------------- 自然高度与分派
NATURAL = {
    'note': lambda item, w: _note_h(item, item.get('size', SIZE['note']), w),
    'markers': lambda item, w: _markers_h(_bullet_items(item['items']), item.get('size', SIZE['markers']),
                                          w - (0.32 if item.get('fill') else 0), item.get('after', 5)) +
    (0.32 if item.get('fill') else 0),
    'section': lambda item, w: _section_h(item, item.get('size', SIZE['markers']), w),
    'chevrons': lambda item, w: _chev_h(item, item.get('size', SIZE['chevrons']), w),
    'para': lambda item, w: _para_h(item, item.get('size', SIZE['para']), w),
    'hexes': lambda item, w: item.get('h', 0.85),
    'tags': lambda item, w: item.get('h', 0.62),
    'pill': lambda item, w: item.get('h', item.get('size', SIZE['pill']) / 72 * 1.9 + 0.02),
    'arrow': lambda item, w: item.get('h', 0.5),
    'process': lambda item, w: item.get('h', 0.48),
    'links': lambda item, w: item.get('h', len(item['items']) * 0.58),
    'label_rows': lambda item, w: _label_rows_h(item, w),
    'callout': lambda item, w: measure(item['text'], 18, w - 0.45, True, spacing=1.2)[1] + 0.3,
    'kv': lambda item, w: item.get('h', 0.44),
}
GROWABLE = {'note', 'markers', 'chevrons', 'callout', 'tags', 'table', 'label_rows', 'section'}


def _info_block(block):
    if getattr(block, '_info', False):
        return block

    def inner(slide, spec):
        kind = spec.get('type')
        if kind in ('figure', 'paper_figure'):
            figure(slide, spec, block)
        elif kind == 'table':
            info_table(slide, spec)
        else:
            block(slide, spec)
    inner._info = True
    inner._raw = block
    return inner


def dispatch(slide, spec, block):
    kind = spec['type']
    block = _info_block(block)
    simple = {'note': note, 'hexes': hexes, 'tags': tags, 'pill': pill, 'markers': markers, 'chevrons': chevrons,
              'arrow': arrow, 'side_tag': side_tag, 'callout': callout, 'layers': layers, 'steps_v': steps_v,
              'quadrant': quadrant, 'para': para, 'process': process, 'links': links, 'label_rows': label_rows,
              'section': section, 'cycle': cycle, 'kv': kv, 'numbered': numbered, 'branch': branch}
    with_block = {'stages': stages, 'timeline': timeline, 'rows': rows, 'fig_list': fig_list, 'card': card,
                  'matrix': matrix}
    if kind in simple:
        simple[kind](slide, spec)
    elif kind in with_block:
        with_block[kind](slide, spec, block)
    else:
        return False
    return True


# ---------------------------------------------------------------- 整页
def render_page(slide, spec, number, meta, block):
    layouts.CURRENT_PAGE = f'第{number}页「{spec.get("title", "")}」' + (f'（{spec["page_type"]}）' if spec.get('page_type') else '')
    _PAGE_BIG[0] = 0
    headline(slide, spec['title'], meta)
    y = LEAD_Y
    if spec.get('lead'):
        y = lead(slide, spec['lead'], y) + GAP_Y
    if spec.get('sublead'):
        y = sublead(slide, spec['sublead'], y - 0.04) + 0.04
    if spec.get('band'):
        b = _norm(spec['band'])
        if b.get('images'):
            hb = b.get('h', 0.6)
            image_band(slide, b, X0, y, W0, hb)
            y += hb + GAP_Y
        else:
            band(slide, b['text'], X0, y, W0)
            y += BAND_H + GAP_Y
    if spec.get('process'):
        pr = spec['process'] if isinstance(spec['process'], dict) else {'items': spec['process']}
        process(slide, dict(pr, box=[X0, y, W0, 0.46]))
        y += 0.46 + GAP_Y
    footer = meta.get('info_footer', False)
    bottom_limit = 7.24 if footer else BOTTOM
    summary_y = bottom_limit - SUMMARY_H if spec.get('summary') else None
    bottom = (summary_y - GAP_Y) if summary_y else bottom_limit
    cb = (X0, y, W0, bottom - y)
    preset = INFO_LAYOUTS.get(spec.get('layout', 'info_full'), {})
    place = layouts.place_factory(slide, _info_block(block))
    for slot, region in spec.get('regions', {}).items():
        default = preset.get(slot)
        default_fill = 'frame'
        if isinstance(default, tuple):
            default, default_fill = default
        g = region.get('grid') or default
        if g is None:
            raise ValueError(f"info 版式 {spec.get('layout')} 没有槽位 '{slot}'，请写 grid；可选：{sorted(preset)}")
        rx, ry, rw, rh = grid_box(cb, g)
        fill = region.get('fill', default_fill)
        if fill in ('frame', 'dashed'):
            rect(slide, rx, ry, rw, rh, 'FFFFFF', LINE, dash=fill == 'dashed', name='分区框-' + slot)
        elif fill in ('note', 'light'):
            rect(slide, rx, ry, rw, rh, LIGHT, name='分区框-' + slot)
        elif fill not in (None, 'none'):
            rect(slide, rx, ry, rw, rh, fill, name='分区框-' + slot)
        xi, wi = rx, rw
        if region.get('side'):
            # 左侧深蓝标签：窄（<0.9 in）时竖排（参考图 2.1），宽时横排两行（模板 2.1 “适用场景与选择建议”）
            sw = region.get('side_w', 0.44)
            rect(slide, rx, ry, sw, rh, NAVY, name='分区竖标签')
            label = region['side'] if sw >= 0.9 else '\n'.join(region['side'].replace('\n', ''))
            tx(slide, label, rx + 0.06, ry, sw - 0.12, rh, region.get('side_size', SIZE['side'] + (2 if sw >= 0.9 else 0)),
               'FFFFFF', True, 'center', spacing=1.05 if sw >= 0.9 else 1.0)
            xi, wi = rx + sw, rw - sw
        top, foot_h = 0.0, 0.0
        if region.get('band'):
            bh = region.get('band_h', 0.46)
            band(slide, region['band'], xi, ry, wi, bh, region.get('band_size', SIZE['region_band']), region.get('num'))
            top = bh
        if region.get('foot'):
            foot_h = 0.4
            rect(slide, xi, ry + rh - foot_h, wi, foot_h, NAVY, name='分区底条')
            tx(slide, region['foot'], xi + 0.1, ry + rh - foot_h, wi - 0.2, foot_h, SIZE['foot'], 'FFFFFF', True,
               'center', name='分区底条文字')
        pad = region.get('pad', 0.1 if fill not in (None, 'none') else 0.0)
        if region.get('head'):
            _head(slide, region['head'], xi + pad, ry + top + pad, wi - 2 * pad, 0.42,
                  region.get('num') if not region.get('band') else None, region.get('head_style', 'fill'))
            top += 0.42 + pad
        if region.get('title'):
            tx(slide, region['title'], xi + pad, ry + top + 0.04, wi - 2 * pad, 0.36, SIZE['head'], NAVY, True,
               'center', name='分区标题')
            top += 0.4
        inner_top = top + (0.08 if (top or pad) else 0)
        avail = rh - inner_top - pad - (foot_h + 0.06 if foot_h else 0)
        items = [dict(it, grow=True) if it['type'] in ('table', 'label_rows') and 'h' not in it else it
                 for it in region.get('items', [])]          # 表格撑满分区，不留半截空白
        if items and region.get('grow') and all(layouts.natural_height(it, wi - 2 * pad) is not None for it in items):
            items = [dict(it, grow=True) if it['type'] in GROWABLE and 'h' not in it else it for it in items]
        layouts.CURRENT_REGION = slot
        layouts.stack(items, xi + pad, ry + inner_top, wi - 2 * pad, avail, place, region.get('gap', 0.1),
                      region.get('valign', 'middle'), f"分区 '{slot}'" if fill not in (None, 'none') else None)
    if spec.get('summary'):
        summary(slide, spec['summary'], summary_y)
    pt = spec.get('page_type')
    if _PAGE_BIG[0] > 1:
        print(f'  ! [{layouts.CURRENT_PAGE}] 一页放了 {_PAGE_BIG[0]} 张大图：一页最多 1 张，其余拆页或降为中图')
    if pt == 'P04' and _PAGE_BIG[0] == 0:
        print(f'  ! [{layouts.CURRENT_PAGE}] P04 大图解读页没有标 role: "big" 的图')
    if footer:
        rect(slide, 0, 7.24, 13.333333, 0.26, NAVY)
        slide_number(slide, number, 12.3, 7.24, 0.73, 0.25)


# ---------------------------------------------------------------- 封面 / 目录 / 结束页（同一套蓝白规范）
def _brand_logo(meta, kind='header'):
    if kind == 'header':
        return meta.get('header_logo') or (str(ASSETS / 'branding' / 'xmu-energy-header.png') if meta.get('brand') == 'xmu' else None)
    return meta.get('logo') or (str(ASSETS / 'branding' / 'xmu-seal.png') if meta.get('brand') == 'xmu' else None)


def _frame_bands(slide):
    rect(slide, 0, 0, 13.333, 0.14, NAVY, name='顶部色带')
    rect(slide, 0, 7.12, 13.333, 0.38, NAVY, name='底部色带')
    rect(slide, 0, 7.04, 13.333, 0.08, ARROW, name='底部浅色带')


def cover_page(slide, spec, meta):
    """封面：顶部 / 底部深蓝色带，左上机构标识，居中题目（40pt，长题分两行）+ 副标题 + 汇报信息行。
    style: panel 时改成左侧深蓝竖栏（tagline + 主题图）+ 右侧左对齐题目（组会、文献、竞选等通用汇报）。"""
    if spec.get('style') == 'panel':
        return _cover_panel(slide, spec, meta)
    _frame_bands(slide)
    logo = _brand_logo(meta)
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(0.6), Inches(0.5), height=Inches(0.62))
        pic.name = '机构标识'
    elif meta.get('institution'):
        tx(slide, meta['institution'], 0.6, 0.5, 6, 0.5, 18, NAVY, True)
    title = spec['title']
    size = spec.get('size', 40 if text_width(title) <= 18 else 34)
    title_box = spec.get('title_box', [0.8, 2.15, 11.73, 1.5])
    tx(slide, title, *title_box, size, NAVY, True, 'center', name='封面题目')
    if spec.get('title_underline', True):
        rect(slide, 6.07, 3.78, 1.2, 0.06, NAVY, name='题目下划线')
    if spec.get('subtitle'):
        tx(slide, spec['subtitle'], 1.0, spec.get('subtitle_y', 3.98), 11.33, 0.5, 20, INK, False, 'center',
           name='副标题')
    info = spec.get('info') or [v for v in (spec.get('author', meta.get('author')), spec.get('advisor', meta.get('advisor')),
                                           spec.get('institution', meta.get('institution')),
                                           spec.get('date', meta.get('date'))) if v]
    if info:
        w = min(3.2, 11.7 / len(info))
        x0 = (13.333 - w * len(info)) / 2
        for i, v in enumerate(info):
            if i:
                rect(slide, x0 + i * w - 0.01, 5.28, 0.02, 0.38, LINE, name='信息分隔')
            tx(slide, v, x0 + i * w, 5.2, w, 0.54, 17, INK, False, 'center', name='汇报信息')


def agenda_page(slide, spec, meta):
    """目录：左侧深蓝竖栏“目录 CONTENTS”，右侧编号条目；已讲部分变灰，当前部分深蓝加粗并带浅蓝底。"""
    items = spec['items']
    if not 2 <= len(items) <= 7:
        raise ValueError('目录支持 2–7 项')
    rect(slide, 0, 0, 4.1, 7.5, NAVY, name='目录竖栏')
    tx(slide, '目录', 0.6, 2.75, 3.0, 1.0, 54, 'FFFFFF', True, 'center')
    tx(slide, 'CONTENTS', 0.6, 3.75, 3.0, 0.5, 20, 'BCD3F6', True, 'center')
    logo = _brand_logo(meta)
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(9.0), Inches(0.35), height=Inches(0.5))
        pic.left = Inches(12.9) - pic.width
    gap = min(0.95, 5.2 / len(items))
    start = 3.75 - gap * len(items) / 2
    active = spec.get('active', 0)
    for i, value in enumerate(items):
        y = start + i * gap
        state = 'now' if i == active else 'done' if i < active else 'next'
        if state == 'now':
            rect(slide, 4.9, y + 0.06, 7.6, gap - 0.12, LIGHT, name='当前部分底')
        circle = NAVY if state != 'done' else 'A9B4C4'
        d = 0.58
        rect(slide, 5.15, y + (gap - d) / 2, d, d, circle, shape=MSO_SHAPE.OVAL, name='目录编号圆')
        tx(slide, f'{i + 1:02d}', 5.15, y + (gap - d) / 2, d, d, 17, 'FFFFFF', True, 'center')
        color = NAVY if state == 'now' else '8A94A6' if state == 'done' else INK
        tx(slide, value, 6.0, y, 6.4, gap, 26, color, True, name='目录项')


def _cover_panel(slide, spec, meta):
    rect(slide, 0, 0, 3.83, 7.5, NAVY, name='封面竖栏')
    tx(slide, spec.get('panel_title', '汇报'), 0.48, 0.62, 2.9, 0.62, 30, 'FFFFFF', True)
    rect(slide, 0.53, 1.42, 2.74, 0.02, 'FFFFFF', name='竖栏细线')
    if spec.get('tagline'):
        tx(slide, spec['tagline'], 0.49, 1.7, 2.85, 1.6, 21.6, 'FFFFFF', True, anchor='top', spacing=1.3)
    fig = spec.get('figure')
    if fig:
        figure(slide, dict(_norm(fig) if isinstance(fig, dict) else {'file': fig}, type='figure', caption='', fit='cover',
                           box=[0.5, 4.0, 2.83, 2.3]), _COVER_BLOCK[0])
    logo = _brand_logo(meta)
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(9.0), Inches(0.5), height=Inches(0.5))
        pic.left = Inches(12.9) - pic.width
    if spec.get('kicker'):
        tx(slide, spec['kicker'], 4.42, 0.93, 6.0, 0.41, 16.2, NAVY, True)
    tx(slide, spec['title'] + ('\n' + spec['subtitle'] if spec.get('subtitle') else ''), 4.4, 2.0, 8.4, 1.8,
       spec.get('size', 32), NAVY, True, spacing=1.2, name='封面题目')
    rect(slide, 4.45, 4.12, 8.2, 0.03, NAVY, name='题目下划线')
    info = spec.get('info') or []
    if info:
        tx(slide, '\n'.join(info), 4.45, 4.45, 8.0, 1.6, 17.4, INK, False, anchor='top', spacing=1.35, name='汇报信息')
    if spec.get('note'):
        tx(slide, spec['note'], 4.45, 6.55, 8.2, 0.4, 13.8, GREY, False, name='封面备注')


_COVER_BLOCK = [None]       # dense_deck 传入的 block，用于封面图


def end_page(slide, spec, meta):
    """结束页：居中主句（默认“感谢聆听”）+ 副句（默认“敬请各位老师批评指正”）。
    写 points: [{title, text}]（最多 3 条）时改成“回到本次汇报的核心”：编号要点 + 底部浅蓝致谢条。"""
    if spec.get('points'):
        rect(slide, 0, 0, 13.333, 0.26, NAVY, name='顶部色带')
        tx(slide, spec.get('heading', '回到本次汇报的核心'), 0.58, 0.8, 12.17, 0.54, 25.8, NAVY, True)
        pts = spec['points']
        step = min(1.28, 4.0 / len(pts))
        for i, pt in enumerate(pts):
            py = 1.9 + i * step
            _circle_num(slide, i + 1, 0.61, py, 0.57)
            tx(slide, pt['title'], 1.43, py - 0.02, 10.4, 0.42, 19.8, NAVY, True)
            tx(slide, pt.get('text', ''), 1.44, py + 0.46, 10.4, 0.47, 17.4, INK, False, highlights=pt.get('highlights'))
        rect(slide, 0.62, 6.13, 12.11, 0.82, LIGHT, name='致谢条')
        tx(slide, spec.get('title', '感谢聆听，欢迎交流'), 0.83, 6.13, 11.68, 0.82, 26.4, NAVY, True, 'center')
        return
    _frame_bands(slide)
    logo = _brand_logo(meta)
    if logo:
        pic = slide.shapes.add_picture(logo, Inches(0.6), Inches(0.5), height=Inches(0.62))
        pic.name = '机构标识'
    tx(slide, spec.get('title', '感谢聆听'), 1.0, 2.55, 11.33, 1.1, 48, NAVY, True, 'center', name='结束主句')
    rect(slide, 6.07, 3.75, 1.2, 0.06, NAVY, name='题目下划线')
    tx(slide, spec.get('subtitle', '敬请各位老师批评指正'), 1.0, 3.95, 11.33, 0.6, 22, INK, False, 'center',
       name='结束副句')
    if spec.get('info'):
        tx(slide, '　　'.join(spec['info']), 1.0, 5.2, 11.33, 0.5, 17, INK, False, 'center', name='汇报信息')


# ---------------------------------------------------------------- 组会 / 文献 / 竞选 / 活动类场景的组件
def kv(slide, spec):
    """浅蓝信息条：items [["数据版本", "v2.3"], ["控制条件", "同一测试集"], …]，键深蓝粗、值黑色，一行排开。"""
    x, y, w, h = spec['box']
    rect(slide, x, y, w, h, spec.get('fill', LIGHT), name='信息条')
    runs = []
    for i, (k, v) in enumerate(spec['items']):
        if i:
            runs.append(('　　', INK, False))
        runs += [(k + '：', NAVY, True)] + runs_from(str(v), spec.get('highlights'), spec.get('strong'))
    size = fit_size(spec, 'markers', lambda s, ww: measure(''.join(r[0] for r in runs), s, ww, True)[1] + 0.1,
                    w - 0.3, h)
    rich(slide, [{'runs': runs}], x + 0.15, y, w - 0.3, h, size, anchor='middle', name='信息条文字')


def numbered(slide, spec):
    """编号议程行（模板“今天的汇报安排”）：items [{title, text}]，深蓝方块编号 + 深蓝粗标题 + 说明，行间细线。"""
    x, y, w, h = spec['box']
    items = spec['items']
    n = len(items)
    rh = h / n
    tw = spec.get('title_w', 0.32) * w
    for i, it in enumerate(_bullet_items(items)):
        ry = y + i * rh
        d = min(0.62, rh * 0.68)
        rect(slide, x + 0.1, ry + (rh - d) / 2, d, d, NAVY, name='议程编号')
        tx(slide, f'{i + 1:02d}', x + 0.1, ry + (rh - d) / 2, d, d, 19, 'FFFFFF', True, 'center')
        tx(slide, it.get('title', ''), x + d + 0.35, ry, tw, rh, spec.get('size', 20), NAVY, True, name='议程标题')
        rich(slide, [{'runs': runs_from(it.get('text', ''), it.get('highlights'), it.get('strong'))}],
             x + d + 0.45 + tw, ry, w - d - 0.55 - tw, rh, spec.get('text_size', 17), anchor='middle', name='议程说明')
        if i < n - 1 or spec.get('last_rule', True):
            rect(slide, x + d + 0.3, ry + rh - 0.01, w - d - 0.4, 0.012, RULE, name='议程分隔线')


def branch(slide, spec):
    """问题拆解树（模板“当前问题与讨论重点”）：root {title, items} 在左，items [{title, text}] 在右，
    主干 + 分支箭头相连；每个分支是白底细框：左侧浅蓝小标题、右侧一句验证办法。"""
    x, y, w, h = spec['box']
    root = _norm(spec['root'])
    items = spec['items']
    n = len(items)
    rw = w * spec.get('root_w', 0.24)
    gap = spec.get('gap', 0.14)
    ih = (h - gap * (n - 1)) / n
    rh = min(h, max(ih * 1.6, 2.0))
    ry = y + (h - rh) / 2
    rect(slide, x, ry, rw, rh, LIGHT, name='根节点')
    paras = [{'runs': [(root.get('title', ''), NAVY, True)], 'size': SIZE['head'] + 2, 'after': 6}]
    for it in _bullet_items(root.get('items', [])):
        if it.get('label'):
            paras.append({'runs': [(it['label'] + '：', NAVY, True)] + runs_from(it.get('text', ''), it.get('highlights')), 'after': 2})
        else:
            paras.append({'runs': runs_from(it['text'], it.get('highlights')), 'after': 2})
    if root.get('text'):
        paras.append({'runs': runs_from(root['text'], root.get('highlights'))})
    rich(slide, paras, x + 0.15, ry, rw - 0.3, rh, SIZE['markers'], anchor='middle', name='根节点文字')
    bx = x + rw + 0.3                       # 主干位置
    ix = bx + 0.3
    mid = ry + rh / 2
    _line = lambda x1, y1, x2, y2: slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                                              Inches(x2), Inches(y2))
    c = _line(x + rw, mid, bx, mid)
    c.line.color.rgb = RGBColor.from_string(NAVY)
    c.line.width = Pt(1.25)
    tops = [y + i * (ih + gap) + ih / 2 for i in range(n)]
    trunk = _line(bx, min(tops + [mid]), bx, max(tops + [mid]))
    trunk.line.color.rgb = RGBColor.from_string(NAVY)
    trunk.line.width = Pt(1.25)
    hw = (w - (ix - x)) * spec.get('head_w', 0.42)
    for i, it in enumerate(items):
        iy = y + i * (ih + gap)
        _arrow_line(slide, bx, iy + ih / 2, ix - 0.02, iy + ih / 2)
        rect(slide, ix, iy, x + w - ix, ih, 'FFFFFF', LINE, name='分支框')
        hh = min(0.42, ih * 0.45)
        rect(slide, ix + 0.1, iy + 0.1, hw, hh, LIGHT, name='分支标题条')
        tx(slide, it['title'], ix + 0.2, iy + 0.1, hw - 0.2, hh, SIZE['head'], NAVY, True, name='分支标题')
        rich(slide, [{'runs': runs_from(it.get('text', ''), it.get('highlights'), it.get('strong'))}],
             ix + hw + 0.3, iy, x + w - ix - hw - 0.45, ih, SIZE['note'], anchor='middle', name='分支说明')
