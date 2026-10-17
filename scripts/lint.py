"""Pre-render checks for generated decks: overflow, orphan lines, small type, collisions.

Estimates use textfit (CJK 1 em, Latin ~0.55 em) and are deliberately conservative;
every warning should be confirmed on the render. Exit code 1 when errors are found.

    python scripts/lint.py --input deck.pptx [--json report.json]
"""
import argparse
import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
import re
from textfit import measure, visible_count, text_width, CLOSING
try:                     # optional: precise word-split detection (pip install jieba)
    import jieba
    jieba.setLogLevel(60)
except ImportError:
    jieba = None


def split_words(paragraph, lines):
    """Return the words cut by automatic line breaks (needs jieba)."""
    if jieba is None or len(lines) < 2:
        return []
    cuts, pos = set(), 0
    for shown in lines[:-1]:
        pos += len(shown)
        cuts.add(pos)
    spans, start = [], 0
    for word in jieba.cut(paragraph):
        spans.append((start, start + len(word), word))
        start += len(word)
    return [w for a, b, w in spans if len(w.strip()) >= 2 and CJK.search(w) and any(a < c < b for c in cuts)]

CJK = re.compile(r'[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef\u201c\u201d\u2018\u2019]')

EMU = 914400
MIN_PT = {'error': 11, 'body': 12}
A14 = '{http://schemas.microsoft.com/office/drawing/2010/main}m'


def inches(v):
    return Emu(v).inches if v is not None else 0.0


def paragraphs(shape):
    last = 16
    for paragraph in shape.text_frame.paragraphs:
        runs = [r for r in paragraph.runs if r.text]
        if not runs:
            # Blank lines still take a line of height (used as paragraph spacers in the reference style).
            yield '', last, False, [], False, None
            continue
        size = max((r.font.size.pt for r in runs if r.font.size), default=18)
        last = size
        bold = any(r.font.bold for r in runs)
        ppr = paragraph._p.pPr
        bu = ppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') if ppr is not None else None
        bullet = ('sub' if bu.get('char') == '–' else True) if bu is not None else False
        spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else None
        yield ''.join(r.text for r in runs), size, bold, runs, bullet, spacing


def overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    w = min(ax+aw, bx+bw) - max(ax, bx)
    h = min(ay+ah, by+bh) - max(ay, by)
    return max(w, 0) * max(h, 0)


def check(path):
    deck = Presentation(path)
    sw, sh = inches(deck.slide_width), inches(deck.slide_height)
    issues = []
    red_pages = []

    def add(level, slide, shape, message):
        issues.append({'level': level, 'slide': slide, 'shape': shape, 'message': message})

    for number, slide in enumerate(deck.slides, 1):
        texts, solids = [], []
        for shape in slide.shapes:
            box = (inches(shape.left), inches(shape.top), inches(shape.width), inches(shape.height))
            name = shape.name
            x, y, w, h = box
            bx, by, bw, bh = x, y, w, h
            if abs(getattr(shape, 'rotation', 0) or 0) % 180 == 90:      # 竖排旋转文字：按旋转后的外框检查
                cx, cy = x + w / 2, y + h / 2
                bx, by, bw, bh = cx - h / 2, cy - w / 2, h, w
            if bx < -0.01 or by < -0.01 or bx + bw > sw + 0.01 or by + bh > sh + 0.01:
                add('error', number, name, f'超出页面边界 ({bx:.2f},{by:.2f},{bw:.2f},{bh:.2f})')
            if shape.has_chart:
                solids.append((name, box))
                chart = shape.chart
                sizes = []
                for axis_name in ('category_axis', 'value_axis'):
                    try:
                        sizes.append(getattr(chart, axis_name).tick_labels.font.size.pt)
                    except Exception:
                        pass
                if sizes and min(sizes) < MIN_PT['body']:
                    add('warn', number, name, f'图表坐标字号 {min(sizes):g}pt，投影时偏小（建议 ≥12pt）')
                continue
            if shape.shape_type is not None and shape.shape_type == 13:   # picture
                solids.append((name, box))
                continue
            if shape.has_table if hasattr(shape, 'has_table') else False:
                solids.append((name, box))
                table = shape.table
                for row in table.rows:
                    for j, cell in enumerate(row.cells):
                        cw = inches(table.columns[j].width) - inches(cell.margin_left) - inches(cell.margin_right)
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size and r.font.size.pt < MIN_PT['error']:
                                    add('error', number, name, f'表格字号 {r.font.size.pt:g}pt 过小')
                            value = ''.join(r.text for r in p.runs)
                            if not value:
                                continue
                            size = max((r.font.size.pt for r in p.runs if r.font.size), default=14)
                            lines = measure(value, size, cw, any(r.font.bold for r in p.runs))[0]
                            if len(lines) > 1 and 0 < visible_count(lines[-1].strip()) <= 2:
                                add('warn', number, name, f'表格单元格孤行：“…{lines[-2][-5:]}⏎{lines[-1].strip()}”，调整列宽 widths 或缩短')
                            cut = split_words(value, lines)
                            if cut:
                                add('warn', number, name, f'表格单元格换行拆开了词语：{"、".join(cut[:3])}')
                continue
            if not shape.has_text_frame:
                continue
            if shape._element.find('.//' + A14) is not None:
                M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
                if any(shape._element.find('.//' + M + t) is not None for t in ('f', 'rad', 'nary')):
                    add('info', number, name, '公式含分式/根式：LibreOffice 预览只显示线性备用文本，宽度以 PowerPoint 渲染为准')
                    texts.append((name, box))
                    continue
                # Formula: compare the longest fallback line with the box width.
                for p in shape.text_frame._txBody.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
                    rpr = p.getparent().find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
                    size = int(rpr.get('sz', 1800)) / 100 if rpr is not None else 18
                    need = text_width(p.text or '') * size / 72 * 0.92
                    if need > w + 0.05:
                        add('warn', number, name, f'公式“{(p.text or "")[:24]}”约需 {need:.1f} in，框宽 {w:.1f} in')
                texts.append((name, box))
                continue
            content = list(paragraphs(shape))
            if not any(runs for _, _, _, runs, _, _ in content):
                continue
            tf = shape.text_frame
            inner_w = w - inches(tf.margin_left) - inches(tf.margin_right)
            total = 0.0
            def is_white(run):
                try:
                    return str(run.font.color.rgb) == 'FFFFFF'
                except AttributeError:
                    return False
            white = all(is_white(r) for _, _, _, runs, _, _ in content for r in runs)
            if not white and not name.startswith('图上说明'):   # 蓝底反白字、压在图上的说明是有意叠放的
                texts.append((name, box))
            missing = [r.text for _, _, _, runs, _, _ in content for r in runs
                       if CJK.search(r.text) and r._r.rPr is not None and r._r.rPr.get('lang') != 'zh-CN']
            if missing:
                add('warn', number, name, f'中文文本缺少 lang=zh-CN（PowerPoint 会按英文规则断行，标点可能单独成行）：“{missing[0][:12]}”')
            for value, size, bold, runs, bullet, spacing in content:
                if size < MIN_PT['error'] and not white:
                    add('error', number, name, f'字号 {size:g}pt 过小：“{value[:16]}”')
                elif size < MIN_PT['body'] and not white and not name.startswith(('页脚注释', '图注', '标注')):
                    add('warn', number, name, f'字号 {size:g}pt 偏小：“{value[:16]}”')
                lines, height = measure(value, size, inner_w, bold, bool(bullet), spacing, bullet == 'sub')
                total += height - 0.04
                em = size / 72 * (1.04 if bold else 1.0)
                usable = inner_w - (0.56 if bullet == 'sub' else 0.28 if bullet else 0)
                widest = max((text_width(l.rstrip().rstrip(''.join(CLOSING))) * em for l in lines if l.strip()), default=0)
                if widest > usable + 0.06 and not white:
                    add('warn', number, name, f'有一个词或数字比框还宽（约 {widest:.2f} in > {usable:.2f} in），渲染时会被硬拆：减小字号或加宽')
                for shown in lines[:-1] if len(lines) > 1 else []:
                    if shown and shown[-1] in CLOSING and text_width(shown) * em > usable + 0.01:
                        add('warn', number, name, f'行末标点“{shown[-1]}”悬挂在框外（“…{shown[-6:]}”），换种说法或加宽，避免不同软件把它挤到下一行')
                        break
                cut = split_words(value, lines)
                if cut:
                    add('warn', number, name, f'自动换行拆开了词语：{"、".join(cut[:3])}（在语义边界手动换行或缩短）')
                elif jieba is None and len(lines) > value.count('\n') + 1 and usable < 4.2 and CJK.search(value) and not white:
                    add('warn', number, name, f'窄栏中文自动换行 {len(lines)} 行，请在渲染图上确认没有拆词（安装 jieba 可自动检查）')
                if len(lines) > 1:
                    last = lines[-1].strip()
                    if 0 < visible_count(last) <= 2:
                        add('warn', number, name, f'孤行：“…{lines[-2][-6:]}⏎{last}”，改写或调整框宽')
                if len(lines) > 1 and size >= 18 and bold and inner_w > 5:
                    add('info', number, name, f'标题折行为 {len(lines)} 行：“{value[:20]}”')
            if total > h + 0.12 and shape.text_frame.word_wrap is not False:
                add('error' if total > h + 0.35 else 'warn', number, name,
                    f'文字可能溢出：估计 {total:.2f} in > 框高 {h:.2f} in')
        # Text boxes colliding with each other or with charts/pictures/tables.
        for i, (na, ba) in enumerate(texts):
            if na.startswith(('标注', '公式')):
                continue
            area = ba[2] * ba[3]
            for nb, bb in texts[i+1:] + solids:
                if nb.startswith('标注'):
                    continue
                inter = overlap(ba, bb)
                if area and inter > 0.12 * min(area, bb[2] * bb[3]):
                    add('warn', number, na, f'与 {nb} 重叠 {inter:.2f} in²')
        # White space: content should fill gray panels and the content band (y 1.45–6.6 in).
        bodies = []
        for shape in slide.shapes:
            bx, by, bw, bh = (inches(shape.left), inches(shape.top), inches(shape.width), inches(shape.height))
            if by < 1.4 or by > 6.6 or shape.name.startswith(('页脚注释',)):
                continue
            bottom = by + bh
            if shape.has_text_frame and shape.text_frame.text.strip() and shape._element.find('.//' + A14) is None:
                est = 0.04
                for value, size, bold, runs, bullet, spacing in paragraphs(shape):
                    est += measure(value, size, bw, bold, bool(bullet), spacing, bullet == 'sub')[1] - 0.04
                anchor = shape.text_frame._txBody.bodyPr.get('anchor')
                bottom = by + (min(bh, (bh + est) / 2) if anchor == 'ctr' else min(bh, est))
            bodies.append((shape, (bx, by, bw, bh), bottom))
        panels = []
        for shape, (bx, by, bw, bh), _ in bodies:
            try:
                gray = shape.name.startswith('分区框') or (shape.shape_type == 1 and shape.fill.type == 1 and str(shape.fill.fore_color.rgb) in ('F2F2F2', 'F2F3F5'))
            except Exception:
                gray = False
            if gray and bw * bh > 2.5 and not shape.text_frame.text.strip():
                panels.append((shape, (bx, by, bw, bh)))
        closing = any(sh.has_text_frame and sh.text_frame.text.strip() and
                      any(key in sh.text_frame.text for key in ('致谢', '谢谢', '汇报结束', '汇报完毕'))
                      for sh in slide.shapes)
        for shape, (px, py, pw, ph) in panels:
            inside = [bottom for other, (bx, by, bw, bh), bottom in bodies
                      if other is not shape and px - 0.02 <= bx + bw / 2 <= px + pw + 0.02
                      and py - 0.02 <= by + bh / 2 <= py + ph + 0.02 and (bw * bh) < pw * ph * 0.95]
            if not inside:
                continue
            free = py + ph - max(inside)
            if free > max(0.7, 0.28 * ph) and not closing:
                add('warn', number, shape.name, f'灰底分区下方留白 {free:.1f} in（占 {free / ph:.0%}）：先放大已有图表，再补配图 / 小图表 / 公式 / 表格，或缩短、合并分区')
        content_page = any(0.85 < inches(sh.top) < 1.05 and abs(inches(sh.left) - 1.08) < 0.1 and sh.has_text_frame and sh.text_frame.text.strip()
                           for sh in slide.shapes)
        if bodies and not panels and content_page and not closing:
            low = max(bottom for _, _, bottom in bodies)
            if low < 5.3:
                add('warn', number, '页面', f'内容只到 y={low:.1f} in，页面下部约 {6.6 - low:.1f} in 空白：补充支撑内容或放大主视觉')
        placeholders = [s.name for s in slide.shapes if s.is_placeholder]
        if placeholders:
            add('info', number, ', '.join(placeholders), '仍有预留图位（成稿前替换或明确标注）')
        # 红色：强调色，一页最多 2 处（含页底结论条），表格里不用红
        spots, table_red = [], False
        for shape in slide.shapes:
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            if any(is_red(r) and r.text.strip() for r in para.runs):
                                table_red = True
                continue
            if not shape.has_text_frame or shape.name == '说明行标签':   # 标签说明行的红色标签是版式的一部分
                continue
            run_red = False          # 跨段连续的红色（一句话手动断行）算 1 处
            for para in shape.text_frame.paragraphs:
                for r in para.runs:
                    if is_red(r) and r.text.strip():
                        if not run_red:
                            spots.append(r.text.strip())
                        run_red = True
                    elif r.text.strip():
                        run_red = False
        info = any(sh.name in ('引导句', '总结条') for sh in slide.shapes)
        if table_red and not info:     # info 页允许表格里标出关键词（参考图 4.2）
            add('warn', number, '表格', '表格里用了红色：表格本身就是对照结构，红色改到正文判断句或去掉')
        limit = 4 if info else 2
        if len(spots) > limit:
            add('warn', number, '红色', f'红色强调 {len(spots)} 处（{"、".join(spots[:4])}…）：一页最多 {limit} 处，只留最该被记住的判断或数字')
        if not info:
            red_pages.append(bool(spots))
        # 成品稿不出现免责、来源追溯与改稿口气
        for shape in slide.shapes:
            if shape.has_text_frame:
                hit = [k for k in DISCLAIMERS if k in shape.text_frame.text]
                if hit:
                    add('error', number, shape.name, f'成品稿里出现免责/改稿字样“{hit[0]}”：删掉，来源口头讲')
    if len(red_pages) >= 5 and sum(red_pages) > 0.4 * len(red_pages):
        add('warn', 0, '全稿', f'{sum(red_pages)}/{len(red_pages)} 页有红色：强调色宜少（一般不超过三成页面），只留有判断的页面')
    # 小节编号连续：同一部分内 1.1、1.2、1.3…不跳号
    last = {}
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and ((0.85 < inches(shape.top) < 1.05 and inches(shape.left) < 1.3) or shape.name == '页标题编号'):
                m = re.match(r'^\s*(\d+)\.(\d+)(\s|$)', shape.text_frame.text)
                if m:
                    major, minor = int(m.group(1)), int(m.group(2))
                    expect = last.get(major, 0) + 1
                    if minor != expect:
                        add('warn', number, '页内标题', f'小节编号 {major}.{minor} 不连续（应为 {major}.{expect}）：删页或合并后要重排编号')
                    last[major] = minor
    # 页码：写死的数字与实际页序不符（删页 / 插页后常见）
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and inches(shape.top) > 7.15 and inches(shape.left) > 11.5:
                xml = shape._element
                has_field = xml.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}fld') is not None
                value = shape.text_frame.text.strip()
                if value.isdigit() and not has_field and int(value) != number:
                    add('warn', number, '页码', f'页脚写死的页码是 {value}，实际是第 {number} 页：改用幻灯片编号字段')
    return issues


DISCLAIMERS = ('AI生成', 'AI 生成', 'AI-generated', '逐字核对', '待补', '此处插入', '依据材料',
               '不虚构', '教学归纳', '截图位', '仅供版式', '不代表本研究')


def is_red(run):
    try:
        return run.font.color is not None and run.font.color.type is not None and str(run.font.color.rgb) in ('C00000', 'DE1515')
    except AttributeError:
        return False


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--json', type=Path)
    args = parser.parse_args()
    issues = check(args.input)
    order = {'error': 0, 'warn': 1, 'info': 2}
    for item in sorted(issues, key=lambda i: (i['slide'], order[i['level']])):
        print(f"[{item['level']:5}] 第{item['slide']}页 {item['shape']}: {item['message']}")
    counts = {k: sum(i['level'] == k for i in issues) for k in order}
    print(f"共 {counts['error']} 个错误、{counts['warn']} 个警告、{counts['info']} 条提示")
    if args.json:
        args.json.write_text(json.dumps(issues, ensure_ascii=False, indent=2), encoding='utf-8')
    sys.exit(1 if counts['error'] else 0)


if __name__ == '__main__':
    main()
