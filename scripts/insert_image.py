"""Fill a reserved picture slot (图片占位-<label>) with an image at the slot's exact geometry.

python-pptx's insert_picture leaves the new picture without its own size because these
slides have no matching layout placeholder; this helper places the picture explicitly
(contain = whole image letterboxed, cover = native centred crop) and removes the slot.

    python scripts/insert_image.py --input deck.pptx --output deck2.pptx --slide 6 \
        --label "状态 A / 方法 A" --image fig.png [--fit contain|cover]
"""
import argparse
from pathlib import Path
from PIL import Image
from pptx import Presentation


def fill(deck, slide_no, label, image, fit='contain'):
    slide = deck.slides[slide_no - 1]
    targets = [s for s in slide.shapes if s.name == '图片占位-' + label]
    if not targets:
        names = [s.name for s in slide.shapes if s.name.startswith('图片占位-')]
        raise SystemExit(f'No slot "{label}" on slide {slide_no}; available: {names}')
    slot = targets[0]
    x, y, w, h = slot.left, slot.top, slot.width, slot.height
    with Image.open(image) as source:
        ratio = source.width / source.height
    if fit == 'cover':
        picture = slide.shapes.add_picture(str(image), x, y, w, h)
        target = w / h
        if ratio > target:
            picture.crop_left = picture.crop_right = (1 - target / ratio) / 2
        elif ratio < target:
            picture.crop_top = picture.crop_bottom = (1 - ratio / target) / 2
    else:
        fw = min(w, int(h * ratio))
        fh = int(fw / ratio)
        picture = slide.shapes.add_picture(str(image), x + (w - fw) // 2, y + (h - fh) // 2, fw, fh)
    picture.name = label
    slot._element.getparent().remove(slot._element)
    return picture


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--slide', type=int, required=True)
    parser.add_argument('--label', required=True)
    parser.add_argument('--image', type=Path, required=True)
    parser.add_argument('--fit', choices=['contain', 'cover'], default='contain')
    args = parser.parse_args()
    deck = Presentation(args.input)
    fill(deck, args.slide, args.label, args.image, args.fit)
    deck.save(args.output)
    print(f'Filled slide {args.slide} slot "{args.label}" -> {args.output}')
