"""Create native editable academic slides in the learned white-and-blue reference style."""
import argparse
import json
import re
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BLUE, DARK, INK = '1D50A2', '17347D', '000000'
PALE, GRAY, ORANGE, RED = 'E3EFF9', 'F2F2F2', 'FCEBDD', 'C00000'   # GRAY = reference bg1 lumMod 95%


def box(slide, x, y, width, height, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(width), Inches(height))
    item.fill.solid()
    item.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        item.line.color.rgb = RGBColor.from_string(line)
        item.line.width = Pt(1)
    else:
        item.line.fill.background()
    style = item._element.find(qn('p:style'))
    if style is not None:
        effect = style.find(qn('a:effectRef'))
        if effect is not None:
            effect.set('idx', '0')
    return item


def zh_run(run_element):
    """Tag a run as Chinese so PowerPoint applies East Asian line breaking (kinsoku) and
    draws ambiguous-width punctuation (“” ‘’ … —) with the East Asian font."""
    props = run_element.get_or_add_rPr() if hasattr(run_element, 'get_or_add_rPr') else run_element
    props.set('lang', 'zh-CN')
    props.set('altLang', 'en-US')
    return props


def text(slide, value, x, y, width, height, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         highlights=None, anchor=None, bullet=False, spacing=None, strong=None, sub=False):
    """Native text box. highlights → red bold, strong → black bold (both inline, exact substrings).
    bullet=True gives the reference's hanging "•" list with 130% line spacing.
    sub=True gives a second-level hanging "–" mark (indent 0.28 in) for子点，配小字号用。"""
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = item.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    if anchor == 'middle':
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    red, black = list(highlights or []), list(strong or [])
    terms = sorted(red + black, key=len, reverse=True)
    for i, line in enumerate(value.split('\n')):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = spacing or (1.3 if bullet else 1.18)
        paragraph.space_after = Pt(3)
        props = paragraph._p.get_or_add_pPr()
        props.set('latinLnBrk', '0')
        props.set('eaLnBrk', '1')
        props.set('hangingPunct', '1')
        if (bullet or sub) and line.strip():
            props.set('marL', str(Inches(0.56 if sub else 0.28)))
            props.set('indent', str(-Inches(0.28)))
            font = OxmlElement('a:buFont')
            font.set('typeface', 'Arial')
            char = OxmlElement('a:buChar')
            char.set('char', '–' if sub else '•')
            props.append(font)
            props.append(char)
        parts = re.split('('+'|'.join(re.escape(t) for t in terms)+')', line) if terms else [line]
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run()
            run.text = part
            run.font.name = 'Arial'
            run.font.size = Pt(size)
            run.font.bold = bold or part in red or part in black
            run.font.color.rgb = RGBColor.from_string(RED if part in red else color)
            rpr = zh_run(run._r)
            ea = OxmlElement('a:ea')
            ea.set('typeface', '微软雅黑')
            rpr.append(ea)
    return item


def bullets(slide, items, x, y, width, height, size=16):
    item = text(slide, '\n'.join(items), x, y, width, height, size)
    for paragraph in item.text_frame.paragraphs:
        props = paragraph._p.get_or_add_pPr()
        props.set('marL', str(Inches(0.2)))
        props.set('indent', str(-Inches(0.2)))
        bullet = OxmlElement('a:buChar')
        bullet.set('char', '•')
        props.append(bullet)
        paragraph.space_after = Pt(10)
    return item


def header(slide, section, section_title, title, page, meta):
    box(slide, 0.25, 0, 0.66, 0.64, BLUE)
    box(slide, 0.25, 0.32, 0.66, 0.6, BLUE, MSO_SHAPE.OVAL)
    text(slide, section, 0.25, 0.27, 0.66, 0.3, 18, 'FFFFFF', True, PP_ALIGN.CENTER)
    text(slide, section_title, 1.08, 0.23, 7.8, 0.46, 24, BLUE, True)
    brand = Path(__file__).resolve().parents[1] / 'assets' / 'branding' / 'xmu-energy-header.png'
    # Institution marks must be requested explicitly (meta.brand or meta.header_logo).
    logo = meta.get('header_logo') or (str(brand) if meta.get('brand') == 'xmu' else None)
    if logo:
        slide.shapes.add_picture(logo, Inches(10.214), Inches(0.152), width=Inches(2.829))
    else:
        text(slide, meta.get('institution', ''), 9.15, 0.25, 3.75, 0.36, 12, BLUE, align=PP_ALIGN.RIGHT)
    box(slide, 1.13, 0.75, 11.91, 0.083, BLUE)   # reference header rule: 0.083 in
    text(slide, title, 1.08, 0.93, 11.7, 0.5, 18, INK, True)
    box(slide, 0, 7.24, 13.333333, 0.26, BLUE)
    slide_number(slide, page, 12.3, 7.24, 0.73, 0.25)


def slide_number(slide, page, x, y, w, h):
    """页码用 PowerPoint 的“幻灯片编号”字段：删页、插页、调顺序后自动重排，
    不会出现写死的旧页码（用户稿删了两页后，后面每页页码都错了 2）。"""
    import uuid
    item = text(slide, '', x, y, w, h, 12, 'FFFFFF', align=PP_ALIGN.RIGHT)
    item.name = '页码'
    paragraph = item.text_frame.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    field = OxmlElement('a:fld')
    field.set('id', '{%s}' % str(uuid.uuid4()).upper())
    field.set('type', 'slidenum')
    props = OxmlElement('a:rPr')
    props.set('lang', 'zh-CN')
    props.set('sz', '1200')
    fill = OxmlElement('a:solidFill')
    color = OxmlElement('a:srgbClr')
    color.set('val', 'FFFFFF')
    fill.append(color)
    props.append(fill)
    latin = OxmlElement('a:latin')
    latin.set('typeface', 'Arial')
    props.append(latin)
    field.append(props)
    value = OxmlElement('a:t')
    value.text = str(page)
    field.append(value)
    end = paragraph._p.find(qn('a:endParaRPr'))
    if end is not None:
        end.addprevious(field)
    else:
        paragraph._p.append(field)
    return item


def strip(slide, value):
    box(slide, 0.65, 6.43, 12.0, 0.57, ORANGE)
    text(slide, value, 0.85, 6.53, 11.55, 0.35, 16, INK, False)   # 页底结论条：黑色 16pt，红色留给句内判断


def cover(slide, spec, meta):
    box(slide, 5.47, 0, 2.39, 1.995, BLUE)
    box(slide, 5.47, 0.8, 2.39, 2.39, BLUE, MSO_SHAPE.OVAL)
    seal = Path(__file__).resolve().parents[1] / 'assets' / 'branding' / 'xmu-seal.png'
    logo = meta.get('logo') or (str(seal) if meta.get('brand') == 'xmu' else None)
    if logo:
        slide.shapes.add_picture(logo, Inches(5.88), Inches(0.94), width=Inches(1.57))
    else:
        text(slide, '研究\n汇报', 5.97, 0.99, 1.4, 1.15, 30, 'FFFFFF', True, PP_ALIGN.CENTER)
    text(slide, spec['title'], 0.65, 3.65, 12.03, 1.1, 40, BLUE, True, PP_ALIGN.CENTER)
    text(slide, spec.get('subtitle', ''), 1.35, 4.78, 10.63, 0.46, 18, INK, align=PP_ALIGN.CENTER)
    entries = [meta.get('author', ''), meta.get('institution', ''), meta.get('date', '')]
    for i, value in enumerate(entries):
        text(slide, value, 1.25+i*3.75, 5.72, 3.45, 0.55, 18, INK, align=PP_ALIGN.CENTER)


def agenda(slide, spec, meta):
    items = spec['items']
    if not 2 <= len(items) <= 7:
        raise ValueError('Agenda supports 2–7 items; use a different layout for longer agendas')
    box(slide, 1.12, 0, 2.2, 2.25, BLUE)
    box(slide, 1.12, 1.1, 2.2, 2.25, BLUE, MSO_SHAPE.OVAL)
    text(slide, '目录', 1.25, 1.02, 1.94, 1.05, 54, 'FFFFFF', True, PP_ALIGN.CENTER)
    text(slide, 'CONTENTS', 1.28, 2.01, 1.88, 0.35, 18, 'FFFFFF', align=PP_ALIGN.CENTER)
    start, gap = 1.8, min(0.86, 4.65/len(items))
    active = spec.get('active', 0)
    for i, value in enumerate(items):
        y = start+i*gap
        # Reference convention: finished chapters gray, current blue, upcoming dark.
        circle, label = (BLUE, BLUE) if i == active else ('A6A6A6', '7F7F7F') if i < active else (DARK, INK)
        box(slide, 5.2, y, 0.6, 0.6, circle, MSO_SHAPE.OVAL)
        text(slide, str(i+1), 5.2, y+0.07, 0.6, 0.45, 24, 'FFFFFF', align=PP_ALIGN.CENTER)
        text(slide, value, 6.15, y+0.03, 6.62, 0.65, 26, label, True)


def context(slide, spec):
    cards = spec['cards']
    if len(cards) != 2:
        raise ValueError('Context layout requires two left cards; add another slide for more content')
    for i, card in enumerate(cards):
        y = 1.72+i*2.1
        box(slide, 0.65, y, 7.55, 1.86, GRAY)
        text(slide, card['title'], 0.92, y+0.15, 7.02, 0.43, 18, BLUE, True)
        bullets(slide, card['items'], 0.94, y+0.72, 6.95, 1.05, 16)
    box(slide, 8.48, 1.72, 4.17, 3.96, PALE)
    text(slide, spec['panel_title'], 8.75, 1.92, 3.64, 0.5, 18, BLUE, True)
    for i, item in enumerate(spec['panel_items']):
        y = 2.62+i*0.83
        box(slide, 8.79, y, 0.34, 0.34, BLUE, MSO_SHAPE.OVAL)
        text(slide, str(i+1), 8.79, y+0.03, 0.34, 0.24, 11, 'FFFFFF', True, PP_ALIGN.CENTER)
        text(slide, item, 9.3, y-0.04, 3.08, 0.62, 15)
    strip(slide, spec['takeaway'])


def method(slide, spec):
    stages = spec['stages']
    if len(stages) != 3:
        raise ValueError('This process layout requires three stages')
    text(slide, spec.get('intro', ''), 0.75, 1.58, 11.85, 0.65, 16)
    for i, stage in enumerate(stages):
        x = 0.75+i*4.15
        box(slide, x, 2.54, 3.5, 0.75, BLUE, MSO_SHAPE.CHEVRON)
        text(slide, stage['title'], x+0.18, 2.73, 3.05, 0.36, 18, 'FFFFFF', True, PP_ALIGN.CENTER)
        box(slide, x, 3.57, 3.5, 2.17, PALE)
        bullets(slide, stage['items'], x+0.25, 3.85, 3.0, 1.56, 16)
    strip(slide, spec['takeaway'])


def add_chart(slide, definition, x):
    text(slide, definition['title'], x, 1.98, 5.65, 0.46, 18, BLUE, True)
    data = CategoryChartData()
    data.categories = definition['categories']
    for series in definition['series']:
        data.add_series(series['name'], series['values'])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(2.6),
                                   Inches(5.6), Inches(3.27), data).chart
    chart.has_title = False
    chart.has_legend = len(definition['series']) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.name = '微软雅黑'
        chart.legend.font.size = Pt(13)
        chart.legend.include_in_layout = False
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = '微软雅黑'
        axis.tick_labels.font.size = Pt(13)
    # Setting python-pptx font.name writes only latin; chart CJK needs a:ea too.
    for props in chart._chartSpace.xpath('.//a:defRPr | .//a:rPr | .//a:endParaRPr'):
        latin = props.find(qn('a:latin'))
        if latin is None:
            latin = OxmlElement('a:latin')
            props.append(latin)
        latin.set('typeface', 'Arial')
        ea = props.find(qn('a:ea'))
        if ea is None:
            ea = OxmlElement('a:ea')
            props.append(ea)
        ea.set('typeface', '微软雅黑')
    palette = [BLUE, '69ACA3', 'ED9D50']
    for index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(palette[index % len(palette)])
        series.format.line.fill.background()
    chart.value_axis.minimum_scale = 0


def results(slide, spec):
    if len(spec['charts']) != 2:
        raise ValueError('The paired-results layout requires two charts')
    text(slide, spec.get('intro', ''), 0.75, 1.53, 11.85, 0.4, 15)
    for chart, x in zip(spec['charts'], (0.85, 6.94)):
        add_chart(slide, chart, x)
    text(slide, spec.get('note', ''), 0.85, 6.08, 11.5, 0.27, 11, '5B6570')
    strip(slide, spec['takeaway'])


def summary(slide, spec):
    if len(spec['items']) != 3:
        raise ValueError('Summary layout requires three findings')
    for i, item in enumerate(spec['items']):
        y = 1.76+i*1.5
        box(slide, 1.55, y, 10.92, 1.22, GRAY, MSO_SHAPE.ROUNDED_RECTANGLE)
        box(slide, 1.1, y+0.36, 0.64, 0.55, BLUE)
        text(slide, str(i+1), 1.1, y+0.47, 0.64, 0.31, 18, 'FFFFFF', True, PP_ALIGN.CENTER)
        text(slide, item['title'], 1.99, y+0.14, 10.0, 0.39, 18, BLUE, True)
        text(slide, item['body'], 1.99, y+0.65, 10.0, 0.43, 16)


def figure(slide, definition, x, y, width, height):
    """Keep image bytes intact; fit with native PPT cropping or letterboxing.

    Optional keys: crop [left, top, right, bottom] as fractions of the source image
    (native, non-destructive crop for taking one sub-panel out of a paper figure),
    frame (default True; False for paper panels on white), tag (e.g. "(a)").
    """
    label = definition.get('label', '示意图')
    framed = definition.get('frame', True)
    fx, fy, fw, fh = x, y, width, height
    if definition.get('file'):
        path = Path(definition['file'])
        with Image.open(path) as source:
            src_w, src_h = source.size
        cl, ct, cr, cb = definition.get('crop', [0, 0, 0, 0])
        ratio = (src_w * (1 - cl - cr)) / (src_h * (1 - ct - cb))
        if definition.get('fit', 'contain') == 'cover':
            target = width / height
            if ratio > target:
                extra = (1 - cl - cr) * (1 - target / ratio) / 2
                cl, cr = cl + extra, cr + extra
            elif ratio < target:
                extra = (1 - ct - cb) * (1 - ratio / target) / 2
                ct, cb = ct + extra, cb + extra
        else:
            # Shrink the slot to the image's own aspect so frame, image and caption hug.
            fw = min(width, height * ratio)
            fh = fw / ratio
            has_caption = definition.get('caption', label) != ''
            group_h = fh + (0.5 if has_caption else 0)
            fx = x + (width - fw) / 2
            fy = y + max(0, (height + (0.58 if has_caption else 0) - group_h) / 2) if definition.get('valign', 'middle') == 'middle' else y
            fy = min(fy, y + height - fh)
        picture = slide.shapes.add_picture(str(path), Inches(fx), Inches(fy), Inches(fw), Inches(fh))
        picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom = cl, ct, cr, cb
        picture.name = label
        if framed:
            outline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(fx), Inches(fy), Inches(fw), Inches(fh))
            outline.fill.background()
            outline.line.color.rgb = RGBColor.from_string('BFC7D5')
            outline.line.width = Pt(0.75)
            style = outline._element.find(qn('p:style'))
            if style is not None:
                outline._element.remove(style)
            outline.name = '图框-' + label
    else:
        # Visible frame + prompt text stay as ordinary shapes (empty placeholders are hidden
        # in slide show/export); a transparent picture placeholder on top takes insert_picture.
        frame = box(slide, x, y, width, height, GRAY, line='A9C0DA')
        frame.name = '图位-' + label
        tf = frame.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        tf.text = '预留图位\n' + label
        for paragraph in tf.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.space_after = Pt(6)
            props = paragraph._p.get_or_add_pPr()
            props.set('marL', '0')
            props.set('indent', '0')
            props.append(OxmlElement('a:buNone'))
            for run in paragraph.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(14)
                run.font.bold = False
                run.font.color.rgb = RGBColor.from_string('687D93')
                ea = OxmlElement('a:ea')
                ea.set('typeface', '微软雅黑')
                zh_run(run._r).append(ea)
        holder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
        holder.fill.background()
        holder.line.fill.background()
        holder.name = '图片占位-' + label
        nv = holder._element.find(qn('p:nvSpPr')).find(qn('p:nvPr'))
        placeholder = OxmlElement('p:ph')
        placeholder.set('type', 'pic')
        placeholder.set('idx', str(1000 + holder.shape_id))
        nv.append(placeholder)
        style = holder._element.find(qn('p:style'))
        if style is not None:
            holder._element.remove(style)
    if definition.get('tag'):
        text(slide, definition['tag'], fx + 0.04, fy + 0.02, 0.5, 0.32, 14, INK, True).name = '标注-' + definition['tag']
    for annotation in definition.get('annotations', []):
        note = text(slide, annotation['text'], fx+annotation['x']*fw, fy+annotation['y']*fh,
                    min(1.8, fw*0.5), 0.34, 12, RED, True)
        note.name = '标注-' + annotation['text'][:12]
    caption = definition.get('caption', label)
    if caption:
        color = {'red': RED, 'blue': BLUE, 'navy': '1E3F99'}.get(definition.get('caption_color'), definition.get('caption_color') or INK)
        text(slide, caption, x, fy+fh+0.08, width, 0.42, definition.get('caption_size', 13), color,
             definition.get('caption_bold', False), align=PP_ALIGN.CENTER).name = '图注-' + label


def visual(slide, spec):
    mode, figures = spec.get('mode', 'side'), spec['figures']
    if mode == 'side':
        if not 1 <= len(figures) <= 2:
            raise ValueError('Side layout supports one or two figures')
        text(slide, spec['body_title'], 0.75, 1.76, 6.16, 0.5, 19, BLUE, True)
        bullets(slide, spec['items'], 0.83, 2.58, 5.93, 3.1, 17)
        height = 3.42 if len(figures) == 1 else 1.62
        for i, definition in enumerate(figures):
            figure(slide, definition, 7.25, 1.8+i*2.15, 5.3, height)
    elif mode == 'hero':
        if len(figures) != 1:
            raise ValueError('Hero layout supports one principal figure')
        figure(slide, figures[0], 0.75, 1.8, 8.4, 3.9)
        text(slide, spec['body_title'], 9.5, 1.85, 3.0, 0.62, 18, BLUE, True)
        bullets(slide, spec['items'], 9.5, 2.8, 3.02, 2.87, 16)
    else:
        raise ValueError('Visual mode must be side or hero')
    strip(slide, spec['takeaway'])


def gallery(slide, spec):
    figures = spec['figures']
    if len(figures) not in (2, 3, 4):
        raise ValueError('Gallery supports two, three, or four panels')
    text(slide, spec.get('intro', ''), 0.75, 1.52, 11.83, 0.48, 15)
    if len(figures) == 2:
        positions = [(0.75, 2.1, 5.7, 3.28), (6.88, 2.1, 5.7, 3.28)]
    elif len(figures) == 3:
        positions = [(0.75+i*4.1, 2.38, 3.63, 2.6) for i in range(3)]
    else:
        positions = [(0.75+i%2*6.13, 1.99+i//2*2.13, 5.7, 1.56) for i in range(4)]
    for definition, position in zip(figures, positions):
        figure(slide, definition, *position)
    strip(slide, spec['takeaway'])


def image_method(slide, spec):
    stages = spec['stages']
    if len(stages) != 3:
        raise ValueError('Image process layout supports three stages')
    for i, stage in enumerate(stages):
        x = 0.75+i*4.1
        figure(slide, stage['figure'], x, 1.82, 3.63, 1.9)
        text(slide, stage['title'], x, 4.38, 3.63, 0.45, 18, BLUE, True)
        bullets(slide, stage['items'], x+0.05, 5.02, 3.5, 0.96, 15)
    strip(slide, spec['takeaway'])


def table_image(slide, spec):
    rows = [spec['columns']] + spec['rows']
    if not 2 <= len(rows) <= 6:
        raise ValueError('Table image layout supports up to five body rows')
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.75), Inches(1.94),
                                    Inches(8.15), Inches(3.84))
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = shape.table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if i == 0 else GRAY if i%2 else 'FFFFFF')
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(14)
                    run.font.bold = i == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else INK)
                    ea = OxmlElement('a:ea')
                    ea.set('typeface', '微软雅黑')
                    zh_run(run._r).append(ea)
    figure(slide, spec['figure'], 9.26, 1.94, 3.29, 3.23)
    strip(slide, spec['takeaway'])


def build(content, output):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = content['title']
    deck.core_properties.author = content.get('meta', {}).get('author', '')
    layouts = {'context': context, 'method': method, 'results': results, 'summary': summary,
               'visual': visual, 'gallery': gallery, 'image_method': image_method, 'table_image': table_image}
    meta = content.get('meta', {})
    for page, spec in enumerate(content['slides'], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        kind = spec['type']
        if kind == 'cover':
            cover(slide, spec, meta)
        elif kind == 'agenda':
            agenda(slide, spec, meta)
        elif kind in layouts:
            header(slide, spec['section'], spec['section_title'], spec['title'], page, meta)
            layouts[kind](slide, spec)
        else:
            raise ValueError(f'Unsupported starter layout: {kind}; create an appropriate native layout')
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    print(f'Created {len(deck.slides)} editable slides -> {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--content', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    content = json.loads(args.content.read_text(encoding='utf-8-sig'))
    def resolve_images(value):
        if isinstance(value, dict):
            for key, item in value.items():
                if key in ('file', 'logo') and item:
                    candidate = Path(item)
                    value[key] = str(candidate if candidate.is_absolute() else args.content.resolve().parent / candidate)
                else:
                    resolve_images(item)
        elif isinstance(value, list):
            for item in value:
                resolve_images(item)
    resolve_images(content)
    build(content, args.output)
