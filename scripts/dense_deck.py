"""Render content-rich academic pages from editable content blocks or slot layouts."""
import argparse
import json
import re
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from styled_deck import PALE, GRAY, ORANGE, box, text, header, cover, agenda, figure, zh_run
BLUE, RED, INK = '073DA5', 'DE1515', '111722'     # 与 info_style 同一套蓝白规范
from layouts import render_regions, render_conclusions
from textfit import text_width
from omml import to_omml, plain_omml
import info_style

MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
A14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
CHART_KINDS = {'line': XL_CHART_TYPE.LINE_MARKERS, 'bar': XL_CHART_TYPE.COLUMN_CLUSTERED,
               'stacked': XL_CHART_TYPE.COLUMN_STACKED, 'barh': XL_CHART_TYPE.BAR_CLUSTERED}


def _font(run_font, size, color=INK, bold=False):
    run_font.size = Pt(size)
    run_font.bold = bold
    run_font.color.rgb = RGBColor.from_string(color)


def chart(slide, spec):
    """Native chart. Options: kind line|bar|stacked|barh, min/max/major, x_title/y_title,
    labels (data labels) + label_format, number_format, size (axis/legend pt, ≥12 advised)."""
    x, y, w, h = spec['box']
    size = spec.get('size', 12)
    title_height = 0.32 if spec.get('title') else 0
    if title_height:
        text(slide, spec['title'], x, y, w, title_height, size+1, BLUE, True, PP_ALIGN.CENTER)
    data = CategoryChartData()
    flip = spec.get('kind') == 'barh'      # bar charts draw bottom-up; flip so the list reads top-down
    data.categories = list(reversed(spec['categories'])) if flip else spec['categories']
    for series in spec['series']:
        data.add_series(series['name'], list(reversed(series['values'])) if flip else series['values'])
    kind = CHART_KINDS[spec.get('kind', 'bar')]
    item = slide.shapes.add_chart(kind, Inches(x), Inches(y+title_height), Inches(w), Inches(h-title_height), data).chart
    item.has_title = False
    item.has_legend = spec.get('legend', len(spec['series']) > 1)
    if item.has_legend:
        item.legend.position = XL_LEGEND_POSITION.BOTTOM
        item.legend.include_in_layout = False
        item.legend.font.size = Pt(size)
        item.legend.font.name = 'Arial'
        item.legend.font.color.rgb = RGBColor.from_string(INK)
    for axis in (item.category_axis, item.value_axis):
        axis.tick_labels.font.name = 'Arial'
        axis.tick_labels.font.size = Pt(size)
        axis.tick_labels.font.color.rgb = RGBColor.from_string(INK)
    skip = OxmlElement('c:tickLblSkip')   # show every category label
    skip.set('val', '1')
    cat_ax = item.category_axis._element
    anchor = cat_ax.find(qn('c:noMultiLvlLbl'))
    if anchor is not None:
        anchor.addprevious(skip)
    else:
        cat_ax.append(skip)
    item.value_axis.has_major_gridlines = True
    item.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string('D9D9D9')
    if 'min' in spec:
        item.value_axis.minimum_scale = spec['min']
    elif spec.get('zero', True):
        item.value_axis.minimum_scale = 0
    if 'max' in spec:
        item.value_axis.maximum_scale = spec['max']
    if 'major' in spec:
        item.value_axis.major_unit = spec['major']
    if spec.get('number_format'):
        item.value_axis.tick_labels.number_format = spec['number_format']
        item.value_axis.tick_labels.number_format_is_linked = False
    for key, axis in (('x_title', item.category_axis), ('y_title', item.value_axis)):
        if spec.get(key):
            axis.has_title = True
            axis.axis_title.text_frame.text = spec[key]
            for run in axis.axis_title.text_frame.paragraphs[0].runs:
                _font(run.font, size)
    if spec.get('labels'):
        plot = item.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = spec.get('label_format', 'General')
        labels.number_format_is_linked = False
        labels.font.size = Pt(size)
        labels.font.color.rgb = RGBColor.from_string(INK)
        if spec.get('kind', 'bar') in ('bar', 'barh'):
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
    for i, series in enumerate(item.series):
        color = spec['series'][i].get('color', [BLUE, RED, '4D89BD', '7F7F7F'][i % 4])
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        series.format.line.color.rgb = RGBColor.from_string(color)
        if spec.get('kind') == 'line':
            series.format.line.width = Pt(2)
            series.smooth = False
            series.marker.size = 5
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.marker.format.line.color.rgb = RGBColor.from_string(color)
    # python-pptx font.name only writes latin; CJK needs a:ea on every run property.
    for props in item._chartSpace.xpath('.//a:defRPr | .//a:rPr | .//a:endParaRPr'):
        props.set('lang', 'zh-CN')
        props.set('altLang', 'en-US')
        for tag, face in [('latin', 'Arial'), ('ea', '微软雅黑')]:
            node = props.find(qn('a:'+tag))
            if node is None:
                node = OxmlElement('a:'+tag)
                props.append(node)
            node.set('typeface', face)


def table(slide, spec):
    """Native table, all cells centred by default. Options: widths, size, align (per-column override), emphasis [[r, c], ...]
    (1-based body row, 0-based column) shown red bold."""
    x, y, w, h = spec['box']
    rows = [spec['columns']] + spec['rows']
    item = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if spec.get('widths'):
        for column, weight in zip(item.columns, spec['widths']):
            column.width = Inches(w * weight / sum(spec['widths']))
    aligns = spec.get('align', ['center'] * len(rows[0]))   # reference style: all cells centred
    emphasis = {tuple(cell) for cell in spec.get('emphasis', [])}
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = item.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if i == 0 else GRAY if i % 2 else 'FFFFFF')
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.055)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if aligns[j] == 'center' else PP_ALIGN.LEFT
                for run in paragraph.runs:
                    hot = (i, j) in emphasis
                    run.font.name = 'Arial'
                    run.font.size = Pt(spec.get('size', 14))
                    run.font.bold = i == 0 or (j == 0 and spec.get('first_col_bold', True)) or hot
                    run.font.color.rgb = RGBColor.from_string(
                        'FFFFFF' if i == 0 else RED if hot else BLUE if j == 0 and spec.get('first_col_bold', True) else INK)
                    ea = OxmlElement('a:ea')
                    ea.set('typeface', '微软雅黑')
                    zh_run(run._r).append(ea)


def diagram(slide, spec):
    """Nodes in relative coordinates. edges [[a, b], ...]; arrows=False draws plain lines;
    edges=[] for parallel (non-sequential) items. Labels are vertically centred."""
    x, y, w, h = spec['box']
    nodes = {node['id']: node for node in spec['nodes']}
    positions = {key: (x+node['x']*w, y+node['y']*h, node['w']*w, node['h']*h)
                 for key, node in nodes.items()}
    for edge in spec.get('edges', []):
        a, b = edge[0], edge[1]
        ax, ay, aw, ah = positions[a]
        bx, by, bw, bh = positions[b]
        if abs(by-ay) > 0.3 and abs(bx-ax) < max(aw, bw):
            start, end = ((ax+aw/2, ay+ah), (bx+bw/2, by)) if by > ay else ((ax+aw/2, ay), (bx+bw/2, by+bh))
        else:
            start, end = ((ax+aw, ay+ah/2), (bx, by+bh/2)) if bx > ax else ((ax, ay+ah/2), (bx+bw, by+bh/2))
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                      Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1]))
        connector.line.color.rgb = RGBColor.from_string(spec.get('edge_color', BLUE))
        connector.line.width = Pt(1.3)
        if spec.get('arrows', True):
            tail = OxmlElement('a:tailEnd')
            tail.set('type', 'triangle')
            connector._element.spPr.get_or_add_ln().append(tail)
    for key, node in nodes.items():
        nx, ny, nw, nh = positions[key]
        shape = MSO_SHAPE.DIAMOND if node.get('decision') else MSO_SHAPE.OVAL if node.get('circle') else MSO_SHAPE.RECTANGLE
        box(slide, nx, ny, nw, nh, node.get('fill', spec.get('fill', 'FFFFFF')), shape, BLUE)
        text(slide, node['label'], nx+0.05, ny, nw-0.1, nh,
             node.get('size', spec.get('size', 14)), node.get('color', INK), node.get('bold', False),
             PP_ALIGN.CENTER, anchor='middle')


def formula(slide, spec):
    """Native Office Math in the standard mc:AlternateContent wrapper with a text fallback,
    so LibreOffice/WPS still show it. Set latex=true to write lines in a LaTeX subset
    (\\frac, _ ^, \\sqrt, \\hat, Greek, \\sum ...); otherwise lines are linear Unicode."""
    x, y, w, h = spec['box']
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    item.name = '公式'
    item.text_frame.word_wrap = True
    body = item._element.find(qn('p:txBody'))
    for old in body.findall(qn('a:p')):
        body.remove(old)
    size = round(spec.get('size', 18) * 100)
    color_value = BLUE if spec.get('blue') else INK
    align = spec.get('align', 'left')

    def run_props(italic=False):
        props = OxmlElement('a:rPr')
        props.set('sz', str(size))
        if italic:
            props.set('i', '1')
        fill = OxmlElement('a:solidFill')
        color = OxmlElement('a:srgbClr')
        color.set('val', color_value)
        fill.append(color)
        props.append(fill)
        latin = OxmlElement('a:latin')
        latin.set('typeface', 'Cambria Math')
        props.append(latin)
        return props

    for line in spec['lines']:
        paragraph = OxmlElement('a:p')
        props = OxmlElement('a:pPr')
        props.set('algn', {'left': 'l', 'center': 'ctr', 'right': 'r'}[align])
        after = OxmlElement('a:spcAft')
        points = OxmlElement('a:spcPts')
        points.set('val', str(spec.get('spacing', 600)))
        after.append(points)
        props.append(after)
        paragraph.append(props)
        build = to_omml if spec.get('latex') else plain_omml
        mathpara, fallback_text = build(line, run_props, align)
        alt = etree.SubElement(paragraph, '{%s}AlternateContent' % MC, nsmap={'mc': MC, 'a14': A14})
        choice = etree.SubElement(alt, '{%s}Choice' % MC)
        choice.set('Requires', 'a14')
        holder = etree.SubElement(choice, '{%s}m' % A14)
        holder.append(mathpara)
        fallback = etree.SubElement(alt, '{%s}Fallback' % MC)
        run = OxmlElement('a:r')
        run.append(run_props(italic=True))
        value = OxmlElement('a:t')
        value.text = fallback_text
        run.append(value)
        fallback.append(run)
        body.append(paragraph)


def flow(slide, spec):
    """Process chain. style "plain" (default): text joined by blue arrows, no boxes — use this;
    style "box": bordered boxes, only for a genuine multi-step procedure and at most once per page."""
    x, y, w, h = spec['box']
    items = [i if isinstance(i, dict) else {'label': i} for i in spec['items']]
    if spec.get('style', 'plain') == 'plain':
        arrow = '  →  ' if spec.get('direction', 'right') == 'right' else '↓'
        sep = arrow if spec.get('arrows', True) else ('　/　' if spec.get('direction', 'right') == 'right' else '')
        labels = [it['label'].replace('\n', ' ') for it in items]
        if spec.get('direction', 'right') == 'right':
            value = sep.join(labels)
        else:
            value = ('\n' + sep + '\n').join(labels) if sep else '\n'.join(labels)
        strong = [it['label'] for it in items if it.get('bold')]
        red = [it['label'] for it in items if it.get('color') == RED]
        box_ = text(slide, value, x, y, w, h, spec.get('size', 15), INK, False, PP_ALIGN.CENTER,
                    highlights=red, anchor='middle', strong=strong, spacing=1.0)
        for paragraph in box_.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip() in ('→', '↓', '/'):
                    run.font.color.rgb = RGBColor.from_string(BLUE)
                    run.font.bold = True
        return
    n = len(items)
    nodes = []
    if spec.get('direction', 'right') == 'right':
        gap = spec.get('gap', 0.32)
        bw = (w - gap * (n - 1)) / n
        bh = min(h, spec.get('box_h', 0.55))
        for i, it in enumerate(items):
            nodes.append(dict(it, id=str(i), x=(i * (bw + gap)) / w, y=(h - bh) / 2 / h, w=bw / w, h=bh / h))
    else:
        gap = spec.get('gap', 0.28)
        bh = min((h - gap * (n - 1)) / n, spec.get('box_h', 0.5))
        total = n * bh + (n - 1) * gap
        top = (h - total) / 2
        for i, it in enumerate(items):
            nodes.append(dict(it, id=str(i), x=0, y=(top + i * (bh + gap)) / h, w=1, h=bh / h))
    edges = [[str(i), str(i + 1)] for i in range(n - 1)] if spec.get('arrows', True) else []
    diagram(slide, {'box': spec['box'], 'nodes': nodes, 'edges': edges, 'size': spec.get('size', 14),
                    'fill': spec.get('fill', 'FFFFFF')})


def stat(slide, spec):
    """已停用：「大数字 + 下面一行小字」的数字块（用户：很丑）。关键数字写进正文句子或表格。"""
    raise ValueError('stat 数字块已停用：把关键数字写进正文句子（可红色强调）或表格')


def shadow(item, blur=90000, dist=25400, direction=5400000, alpha=34000):
    """给自选图形加一层柔和的外阴影（用户规则：并列关系用方框带阴影）。
    PowerPoint 与 LibreOffice 都能渲染 a:outerShdw。"""
    spPr = item._element.spPr
    for old in spPr.findall(qn('a:effectLst')):
        spPr.remove(old)
    effects = OxmlElement('a:effectLst')
    outer = OxmlElement('a:outerShdw')
    outer.set('blurRad', str(blur))
    outer.set('dist', str(dist))
    outer.set('dir', str(direction))
    outer.set('rotWithShape', '0')
    color = OxmlElement('a:srgbClr')
    color.set('val', '000000')
    level = OxmlElement('a:alpha')
    level.set('val', str(alpha))
    color.append(level)
    outer.append(color)
    effects.append(outer)
    spPr.append(effects)


def centered_text(slide, title, body, x, y, w, h, spec, highlights=()):
    """卡片里的文字：标题 + 正文都水平居中，整块在框内垂直居中。"""
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(0)
    rows = [(title, spec.get('title_size', 16), BLUE, True, ()),
            (body, spec.get('size', 15), INK, False, tuple(highlights))]
    first = True
    for value, size, color, bold, reds in rows:
        if not value:
            continue
        terms = sorted(reds, key=len, reverse=True)
        for line in value.split('\n'):
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.line_spacing = 1.28
            paragraph.space_after = Pt(spec.get('space_after', 4))
            parts = re.split('(' + '|'.join(re.escape(t) for t in terms) + ')', line) if terms else [line]
            for part in parts:
                if not part:
                    continue
                run = paragraph.add_run()
                run.text = part
                run.font.name = 'Arial'
                run.font.size = Pt(size)
                run.font.bold = bold or part in terms
                run.font.color.rgb = RGBColor.from_string(RED if part in terms else color)
                props = zh_run(run._r)
                ea = OxmlElement('a:ea')
                ea.set('typeface', '微软雅黑')
                props.append(ea)


def card_height(item, width, spec):
    from textfit import measure
    pad = spec.get('pad', 0.18)
    hgt = 0.0
    if item.get('title'):
        hgt += measure(item['title'], spec.get('title_size', 16), width - 2 * pad, True)[1]
    if item.get('body'):
        hgt += measure(item['body'], spec.get('size', 15), width - 2 * pad, False, spacing=1.28)[1]
    return hgt


def cards(slide, spec):
    """并列卡片：浅蓝同色系圆角 + 柔和阴影，标题与正文都在框内居中（用户规则）。
    items: [{title, body, highlights}]；direction right|down；fills 给一组同色系浅蓝（按顺序取）。"""
    x, y, w, h = spec['box']
    items = spec['items']
    gap = spec.get('gap', 0.25)
    down = spec.get('direction', 'right') == 'down'
    count = len(items)
    fills = spec.get('fills') or [spec.get('fill', 'E7F0FA')]
    line = spec.get('line', 'B9CDE4')
    if not down:
        # 卡片高度按内容定（标题 + 正文行数），不把整块区域撑满；整排在区域内垂直居中
        cw0 = (w - gap * (count - 1)) / count
        need = max(card_height(it, cw0, spec) for it in items)
        ch_row = min(h, spec.get('h_card', need + 0.6))
        y = y + (h - ch_row) / 2
        h = ch_row
    for index, item in enumerate(items):
        if down:
            cw, ch = w, (h - gap * (count - 1)) / count
            cx, cy = x, y + index * (ch + gap)
        else:
            cw, ch = (w - gap * (count - 1)) / count, h
            cx, cy = x + index * (cw + gap), y
        shape = box(slide, cx, cy, cw, ch, fills[index % len(fills)],
                    MSO_SHAPE.ROUNDED_RECTANGLE, line=line)
        try:
            shape.adjustments[0] = spec.get('radius', 0.06)
        except Exception:
            pass
        shadow(shape, alpha=spec.get('shadow_alpha', 32000))
        pad = spec.get('pad', 0.18)
        centered_text(slide, item.get('title', ''), item.get('body', ''),
                      cx + pad, cy + pad * 0.6, cw - 2 * pad, ch - 1.2 * pad, spec, item.get('highlights', []))


BAND_H = 0.44   # 蓝底白粗字标题条的高度（英寸）


def panel(slide, spec):
    """分区。title_style:
       plain（默认）— 蓝色加粗居中小标题；
       band — 整宽蓝底白粗字标题条（参考稿“可靠性评估”一类分区头），并列的 2–4 个模块用；
       tab  — 左上蓝底白粗字标签（参考稿“能源传输设备”一类），标签宽度随文字。"""
    x, y, w, h = spec['box']
    box(slide, x, y, w, h, spec.get('fill', GRAY), line=spec.get('line'))
    title, style = spec.get('title'), spec.get('title_style', 'plain')
    if not title:
        return
    size = spec.get('size', 16)
    if style == 'band':
        box(slide, x, y, w, BAND_H, BLUE).name = '标题条-' + title[:10]
        text(slide, title, x + 0.1, y, w - 0.2, BAND_H, size, 'FFFFFF', True, PP_ALIGN.CENTER, anchor='middle')
    elif style == 'tab':
        tw = min(w, text_width(title) * size / 72 * 1.04 + 0.5)
        box(slide, x, y, tw, BAND_H, BLUE).name = '标签-' + title[:10]
        text(slide, title, x + 0.1, y, tw - 0.2, BAND_H, size, 'FFFFFF', True, PP_ALIGN.CENTER, anchor='middle')
    else:
        # 分区小标题一律居中（用户规则：居中了才好看）
        text(slide, title, x+0.12, y+0.06, w-0.24, 0.4, size, BLUE, True, PP_ALIGN.CENTER)


def transition(slide, spec):
    """参考稿的“A ➜[蓝底白粗字箭头] B”：说明从 A 到 B 的变化及原因。
    spec: left, label（≤6 字）, right, size（默认 16）。整行垂直居中。"""
    x, y, w, h = spec['box']
    size = spec.get('size', 16)
    label = spec['label']
    aw = max(1.35, text_width(label) * 15 / 72 * 1.04 + 0.75)
    ah = min(h, 0.62)
    side = (w - aw - 0.3) / 2
    text(slide, spec['left'], x, y, side, h, size, INK, False, PP_ALIGN.RIGHT, anchor='middle')
    arrow = box(slide, x + side + 0.15, y + (h - ah) / 2, aw, ah, BLUE, MSO_SHAPE.RIGHT_ARROW)
    arrow.name = '转折箭头-' + label
    try:                              # 粗箭身、短箭头，白字放得下（参考稿第 4 页的样子）
        arrow.adjustments[0] = 0.74
        arrow.adjustments[1] = 0.42
    except Exception:
        pass
    text(slide, label, x + side + 0.2, y + (h - ah) / 2, aw - 0.45, ah, 15, 'FFFFFF', True,
         PP_ALIGN.CENTER, anchor='middle')
    text(slide, spec['right'], x + side + aw + 0.3, y, side, h, size, INK, False, PP_ALIGN.LEFT, anchor='middle')


WARNINGS = []


def check_terms(spec, where):
    """高亮/加粗词必须是正文的实际子串，否则静默失效（踩过）。"""
    body = spec.get('text') or spec.get('takeaway') or ''
    for key in ('highlights', 'strong', 'emphasis'):
        for term in spec.get(key) or []:
            if isinstance(term, str) and body and term not in body:
                WARNINGS.append(f'{where}: {key} 词“{term}”不在正文里，会静默失效')


def block(slide, spec):
    x, y, w, h = spec['box']
    kind = spec['type']
    if kind == 'text':
        check_terms(spec, '第%s页' % getattr(slide, '_page_no', '?'))
    if kind == 'panel':
        panel(slide, spec)
    elif kind == 'text':
        text(slide, spec['text'], x, y, w, h, spec.get('size', 16), spec.get('color', INK),
             spec.get('bold', False), PP_ALIGN.CENTER if spec.get('align') == 'center' else PP_ALIGN.LEFT,
             spec.get('highlights'), bullet=spec.get('bullet', False), spacing=spec.get('spacing'),
             strong=spec.get('strong'), sub=spec.get('sub', False))
    elif kind == 'chart':
        chart(slide, spec)
    elif kind == 'table':
        table(slide, spec)
    elif kind == 'diagram':
        diagram(slide, spec)
    elif kind == 'flow':
        flow(slide, spec)
    elif kind == 'stat':
        stat(slide, spec)
    elif kind == 'cards':
        cards(slide, spec)
    elif kind == 'transition':
        transition(slide, spec)
    elif info_style.dispatch(slide, spec, block):
        pass
    elif kind == 'formula':
        formula(slide, spec)
    elif kind == 'figure':
        figure(slide, spec, x, y, w, h)
    elif kind == 'paper_figure':
        # Original paper panel: no gray frame, aspect kept, optional native crop + (a) tag.
        figure(slide, dict({'frame': False, 'fit': 'contain'}, **spec), x, y, w, h)
    else:
        raise ValueError(f'Unsupported block: {kind}')


def _resolve_file(v, base):
    """"@/components/x.png" 指本技能 assets 目录下的内置素材；其余相对路径按内容 JSON 所在目录解析。"""
    if str(v).startswith('@/'):
        return str(Path(__file__).resolve().parents[1] / 'assets' / str(v)[2:])
    return str(Path(v) if Path(v).is_absolute() else base / v)


def resolve_files(value, base):
    if isinstance(value, dict):
        return {k: (_resolve_file(v, base) if k == 'file' and v else resolve_files(v, base))
                for k, v in value.items()}
    if isinstance(value, list):
        return [resolve_files(v, base) for v in value]
    return value


def build(content, output, base):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = content['title']
    meta = dict(content.get('meta', {}))
    for key in ('header_logo', 'logo'):
        if meta.get(key):
            file = Path(meta[key])
            meta[key] = str(file if file.is_absolute() else base/file)
    for number, spec in enumerate(content['slides'], 1):
        spec = resolve_files(spec, base)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide._page_no = number
        if spec.get('type') == 'cover':
            info_style._COVER_BLOCK[0] = block
            info_style.cover_page(slide, spec, meta)
        elif spec.get('type') == 'agenda':
            info_style.agenda_page(slide, spec, meta)
        elif spec.get('type') == 'end':
            info_style.end_page(slide, spec, meta)
        else:
            for key in ('lead', 'summary'):
                if isinstance(spec.get(key), dict):
                    check_terms({'takeaway': spec[key]['text'], 'emphasis': spec[key].get('highlights')}, f'第{number}页{key}')
            info_style.render_page(slide, spec, number, meta, block)
        # Speaker notes are for the speaker only; layout metadata stays in the content JSON.
        if spec.get('notes'):
            slide.notes_slide.notes_text_frame.text = spec['notes']
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    pending = []

    def collect(value, page):
        if isinstance(value, dict):
            if value.get('type') in ('figure', 'paper_figure') and value.get('prompt') and not value.get('file'):
                pending.append((page, value.get('label', 'figure'), value['prompt'], value.get('aspect', '')))
            for item in value.values():
                collect(item, page)
        elif isinstance(value, list):
            for item in value:
                collect(item, page)
    for number, spec in enumerate(content['slides'], 1):
        collect(spec, number)
    if pending:
        folder = output.parent / 'images'
        folder.mkdir(exist_ok=True)
        print('待生成配图（生成后把图片路径写回该图位的 file，再重新生成）：')
        for page, label, prompt, aspect in pending:
            name = ''.join(c if c.isalnum() else '_' for c in label)[:40]
            (folder / f'{name}.prompt.md').write_text(f'# 第{page}页 · {label} {aspect}\n\n{prompt}\n', encoding='utf-8')
            print(f'  第{page}页 {label} -> images/{name}.prompt.md')
    deck.save(output)
    for warning in WARNINGS:
        print('  ! ' + warning)
    print(f'Built {len(deck.slides)} pages -> {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--content', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(json.loads(args.content.read_text(encoding='utf-8-sig')), args.output, args.content.resolve().parent)
