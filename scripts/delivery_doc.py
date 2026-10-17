"""把「交付说明」（markdown）转成可直接转发的 docx。

用法：
    python scripts/delivery_doc.py --input 交付说明.md --output out/交付说明.docx

支持标题 / 段落 / 表格 / 有序无序列表 / **加粗**——交付说明只用得到这几种元素，
所以不必为转个格式去装 pandoc（Homebrew 在本机不可用）。
中文字体走 w:eastAsia=微软雅黑，否则 Word 里会是宋体。
"""
import argparse
import pathlib
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = '1D50A2'
BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
BULLET_RE = re.compile(r'^\s*[-*] ')
NUMBER_RE = re.compile(r'^\s*\d+\. ')


def add_runs(paragraph, text):
    """把 **加粗** 拆成真 run，其余按普通文本。"""
    pos = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_separator(line):
    return '|' in line and set(line.replace('|', '').strip()) <= set('-: ')


def is_table_row(line):
    return line.lstrip().startswith('|')


def split_row(line):
    return [cell.strip() for cell in line.strip().strip('|').split('|')]


def convert(source, target):
    lines = source.read_text(encoding='utf-8').split('\n')
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        # 表格：表头行 + 分隔行 + 若干数据行
        if is_table_row(line) and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            header = split_row(line)
            body, cursor = [], index + 2
            while cursor < len(lines) and is_table_row(lines[cursor]):
                body.append(split_row(lines[cursor]))
                cursor += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = 'Light Grid Accent 1'
            for column, value in enumerate(header):
                cell = table.rows[0].cells[column]
                cell.text = ''
                add_runs(cell.paragraphs[0], value)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9)
            for row in body:
                cells = table.add_row().cells
                for column, value in enumerate(row[:len(header)]):
                    cells[column].text = ''
                    add_runs(cells[column].paragraphs[0], value)
                    for run in cells[column].paragraphs[0].runs:
                        run.font.size = Pt(9)
            doc.add_paragraph()
            index = cursor
            continue
        if line.startswith('### '):
            heading = doc.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            heading = doc.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            heading = doc.add_heading(line[2:], level=1)
        else:
            heading = None
        if heading is not None:
            for run in heading.runs:
                run.font.color.rgb = RGBColor.from_string(BLUE)
        elif line.strip() == '---':
            doc.add_paragraph()
        elif BULLET_RE.match(line):
            add_runs(doc.add_paragraph(style='List Bullet'), BULLET_RE.sub('', line))
        elif NUMBER_RE.match(line):
            add_runs(doc.add_paragraph(style='List Number'), NUMBER_RE.sub('', line))
        else:
            add_runs(doc.add_paragraph(), line)
        index += 1

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return len(doc.paragraphs), len(doc.tables)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=pathlib.Path, required=True)
    parser.add_argument('--output', type=pathlib.Path, required=True)
    args = parser.parse_args()
    paragraphs, tables = convert(args.input, args.output)
    print(f'Wrote {args.output} ({paragraphs} paragraphs, {tables} tables)')


if __name__ == '__main__':
    main()
