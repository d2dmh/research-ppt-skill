#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 mp4 真正嵌入 PPTX —— 不是做链接、不是放截图。

做法（与 PowerPoint 自己写出的结构一致）：
  1) 页面上的封面图仍作静帧：p:pic/p:blipFill 指向封面 png；
  2) 视频作为媒体部件放进 ppt/media/；
  3) 在 p:nvPr 里加 <a:videoFile r:link=…/> 与 <p14:media r:embed=…/>（两条关系：video + media）；
  4) cNvPr 里加 <a:hlinkClick r:id="" action="ppaction://media"/>，点击封面即播放。

用法：
  embed_video.py --input in.pptx --output out.pptx \
                 --spec "15:images/录屏封面-案例一.png:../录屏/案例一.mp4" \
                 --spec "20:images/录屏封面-案例二.png:../录屏/案例三.mp4"
  第一个数字是页码（1 起）；脚本自动挑选该页最大的图片形状作为视频载体。
"""
import argparse, os, shutil, zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
from pptx.opc.packuri import PackURI
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main',
}


def embed(slide, video_path, picture_shape):
    """把 picture_shape 变成内嵌视频的载体。"""
    package = slide.part.package
    base = os.path.basename(video_path)
    partname = PackURI('/ppt/media/' + base)
    blob = open(video_path, 'rb').read()
    media_part = Part(partname, 'video/mp4', package, blob)
    r_video = slide.part.relate_to(media_part, RT.VIDEO)
    r_media = slide.part.relate_to(media_part, RT.MEDIA)

    pic = picture_shape._element
    nv_pr = pic.find('./p:nvPicPr/p:nvPr', NS)
    c_nv_pr = pic.find('./p:nvPicPr/p:cNvPr', NS)
    # 名称改成视频文件名，PowerPoint 里媒体对象会显示这个
    c_nv_pr.set('name', base)
    # 清掉此前挂过的“文件超链接”——同一个 cNvPr 里只能有一个 hlinkClick，否则 PowerPoint 报修复
    for old in c_nv_pr.findall('{%s}hlinkClick' % NS['a']):
        c_nv_pr.remove(old)
    # 点击播放
    hlink = etree.Element('{%s}hlinkClick' % NS['a'])
    hlink.set('{%s}id' % NS['r'], '')
    hlink.set('action', 'ppaction://media')
    c_nv_pr.insert(0, hlink)
    # a:videoFile 必须放在 nvPr 最前，p:extLst 放最后
    vf = etree.Element('{%s}videoFile' % NS['a'])
    vf.set('{%s}link' % NS['r'], r_video)
    nv_pr.insert(0, vf)
    ext_lst = etree.SubElement(nv_pr, '{%s}extLst' % NS['p'])
    ext = etree.SubElement(ext_lst, '{%s}ext' % NS['p'])
    ext.set('uri', '{DAA4B4D4-6D71-4841-9C94-04AC343B8D0A}')
    media = etree.SubElement(ext, '{%s}media' % NS['p14'], nsmap={'p14': NS['p14']})
    media.set('{%s}embed' % NS['r'], r_media)
    return base, len(blob)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True)
    ap.add_argument('--output', required=True)
    ap.add_argument('--spec', action='append', required=True,
                    help='页码:封面图:视频文件')
    a = ap.parse_args()

    prs = Presentation(a.input)
    for spec in a.spec:
        page, poster, video = spec.split(':', 2)
        slide = prs.slides[int(page) - 1]
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        if not pics:
            raise SystemExit('第 %s 页没有图片形状' % page)
        target = max(pics, key=lambda sh: Emu(sh.width).inches if sh.width else 0)
        if Emu(target.width).inches < 3:
            raise SystemExit('第 %s 页找不到大图（最大才 %.2f in）' % (page, Emu(target.width).inches))
        name, size = embed(slide, video, target)
        print('P%-3s 已嵌入视频 %-34s %5.1f MB（载体图 %.2f×%.2f in）'
              % (page, name, size / 1048576, Emu(target.width).inches, Emu(target.height).inches))

    prs.save(a.output)
    # 校验内容类型
    with zipfile.ZipFile(a.output) as z:
        ct = z.read('[Content_Types].xml').decode('utf-8')
        media = [n for n in z.namelist() if n.startswith('ppt/media/')]
        vids = [m for m in media if m.lower().endswith('.mp4')]
        print('内容类型含 mp4 默认项:', 'Extension="mp4"' in ct)
        print('包内视频部件:', vids)
        rels = z.read('ppt/slides/_rels/slide%d.xml.rels' % int(a.spec[0].split(':')[0])).decode('utf-8')
        print('该页关系含 video:', 'relationships/video' in rels)


if __name__ == '__main__':
    main()
